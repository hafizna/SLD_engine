"""/ingest -- stateless: upload -> draft blob -> review -> preview -> publish.

Nothing touches the DB until publish. The draft is a JSON blob the client
re-sends on every call.
"""
import importlib
import json
import os
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

SAMPLE_JSON = Path(__file__).resolve().parent.parent / "samples" / "ss_cwd_ingest.json"
SAMPLE_XLSX = Path(__file__).resolve().parent.parent / "samples" / "ss_cwd_ingest.xlsx"
LBK_XLSX = Path(__file__).resolve().parent.parent / "samples" / "ss_lbk_ingest.xlsx"
BALI_JSON = Path(__file__).resolve().parent.parent / "samples" / "ss_bali_ingest.json"


def test_system_ibt_workbook_keeps_table_12_separate_from_transmission_risks():
    from app.services.ingest_parser import parse_upload
    path = SAMPLE_XLSX.parent / 'system_ibt_500_ingest.xlsx'
    payload = parse_upload(path.read_bytes(), path.name)
    assert payload['meta']['analytical_hint'] == 'IBT_500_150'
    assert not payload['meta']['dropped_edges']
    assert [r['seq_no'] for r in payload['risks']] == list(range(1,39))
    assert all(r['pin_kind'] == 'SUBSTATION' and r['pin_key'] for r in payload['risks'])
    assert payload['risks'][18]['category'] == 'N-1-1'
    assert payload['risks'][1]['category'] == 'BELUM_DITETAPKAN'


def test_muarakarang_reviewed_ibt_and_cross_view_bays():
    from app.services.ingest_parser import parse_upload
    path = SAMPLE_XLSX.parent / 'ss_muarakarang_durikosambi_ingest.xlsx'
    payload = parse_upload(path.read_bytes(), path.name)
    nodes = {n['external_key']: n for n in payload['objects']}
    assert {k for k, n in nodes.items() if n['object_type'] == 'GITET'} == {'GITET_MKBRU', 'GITET_DKSBI'}
    links = {(e['from_external_key'], e['to_external_key'], e['unit_no'])
             for e in payload['connections'] if e['relation_type'] == 'IBT_LINK'}
    assert links == {('GITET_MKBRU', 'GIS MKBRU', '1'),
                     ('GITET_MKBRU', 'GIS MKBRU', '2'), ('GITET_DKSBI', 'DKSBI', '1')}
    assert nodes['DMGOT']['object_type'] == 'GIS'
    assert nodes['DMGOT']['bay_feeder_key'] == 'PINKA'
    assert nodes['DMGOT']['bay_view_keys'] == ['MUARAKARANG']
    assert nodes['PINKA']['bay_feeder_key'] == 'DMGOT'
    assert nodes['PINKA']['bay_view_keys'] == ['DURIKOSAMBI']
    pairs = {frozenset((e['from_external_key'], e['to_external_key'])) for e in payload['connections']}
    assert frozenset(('DMGOT', 'PINKA')) in pairs
    assert frozenset(('GIS MKBRU', 'MKBRU')) in pairs
    assert frozenset(('KBJRK', 'PINKA')) not in pairs


def test_muarakarang_both_views_render_without_geometry_errors():
    from scripts.audit_sample_workbooks import audit_one
    result = audit_one(SAMPLE_XLSX.parent / 'ss_muarakarang_durikosambi_ingest.xlsx')
    assert result['ok'], result
    assert len(result['views']) == 2


def test_step_down_bus_feeding_a_lower_tier_routes_clean():
    """Bengkulu's Pekalongan 70 sits level with its 150 kV bus in the deck and
    still feeds Sukamerindu on Tier-3. The router used to find no channel for
    that line, and a fractional bus y printed 0.1 px off its wires' ends."""
    from scripts.audit_sample_workbooks import audit_one
    path = SAMPLE_XLSX.parent / 'ss_bengkulu_ingest.xlsx'
    result = audit_one(path)
    assert result['ok'], result

    from app.services.ingest_parser import parse_upload
    payload = parse_upload(path.read_bytes(), path.name)
    by_seq = {r['seq_no']: r for r in payload['risks']}
    # Multi Pin = Ya: risk #1 names two ruas, #3 three, #4 two ruas + two GIs
    assert len(by_seq[1]['extra_pins']) == 1
    assert len(by_seq[3]['extra_pins']) == 2
    assert len(by_seq[4]['extra_pins']) == 3


def test_sumsel_draws_its_ibt_only_source_and_routes_past_offset_step_downs(client):
    """Two things Sumsel exposed. Sungai Lilin's 150 kV bus is fed only by its
    IBT (its SUTET lives on the backbone sheet); it used to be filed as a 'bay'
    of its own GITET and both vanished. And an offset 150/70 chain's obstacle
    ran from the HV row down, sealing the exit of PLTU Sumbagsel-1, which sits
    above the pushed Palembang 70 kV bus -- the router found no channel."""
    path = SAMPLE_XLSX.parent / 'ss_sumsel_ingest.xlsx'
    r = client.post("/api/ingest/parse-file", files={"file": (path.name, path.read_bytes(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    draft = r.json()
    for n in draft["nodes"]:
        n["resolution"] = "NEW"
    published = client.post('/api/ingest/publish', json={
        'draft': draft, 'subsystem_code': 'SS_SUMSEL', 'subsystem_name': 'Sumsel'})
    assert published.status_code == 200, published.text
    view = next(v for v in client.get('/api/views').json() if v['view_key'] == 'SS_SUMSEL_FULL')
    svg = client.get(f"/api/views/{view['id']}/sld.svg").text
    assert 'data-code="SGLIN"' in svg and 'data-code="SGLIN_275"' in svg
    assert 'TIDAK masuk gambar utama' not in svg
    from tests.test_sld_geometry import geometry_errors
    assert geometry_errors(svg) == []


def test_sumbagteng_stacks_500_275_150_and_lifts_the_sutet_arrows(client):
    """New Aurduri and Perawang step 500 -> 275 -> 150 kV through two IBTs.
    The GITET placer ran one pass, so the 500 kV bar -- whose LV side is itself
    a GITET -- never got a position and the view failed to render. And a bay on
    a GITET (the SUTET arrow to Muara Enim) hung down into the IBT chain; it
    now leaves from the top of the bar."""
    import re
    path = SAMPLE_XLSX.parent / 'ss_sumbagteng_ingest.xlsx'
    r = client.post("/api/ingest/parse-file", files={"file": (path.name, path.read_bytes(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    draft = r.json()
    for n in draft["nodes"]:
        n["resolution"] = "NEW"
    published = client.post('/api/ingest/publish', json={
        'draft': draft, 'subsystem_code': 'SS_SUMBAGTENG', 'subsystem_name': 'Sumbagteng'})
    assert published.status_code == 200, published.text
    from tests.test_sld_geometry import geometry_errors
    views = {v['view_key']: v for v in client.get('/api/views').json()}
    for side in ('RIAU', 'SUMBAR', 'JAMBI'):
        svg = client.get(f"/api/views/{views[f'SS_SUMBAGTENG_{side}']['id']}/sld.svg").text
        assert geometry_errors(svg) == [], side
    svg = client.get(f"/api/views/{views['SS_SUMBAGTENG_JAMBI']['id']}/sld.svg").text

    def bus_y(code):
        m = re.search(rf'class="sld-node" data-node-kind="SUBSTATION" data-node-id="\d+" '
                      rf'data-code="{code}" data-x="[\d.]+" data-y="([\d.]+)"', svg)
        assert m, code
        return float(m.group(1))
    assert bus_y('NAURD_500') < bus_y('NAURD_275') < bus_y('NAURD')
    arrow = re.search(r'class="sld-bay"[^>]*data-code="MENIM_500".*?d="M[\d.]+,([\d.]+) V([\d.]+)"',
                      svg, re.S)
    assert arrow and float(arrow.group(2)) < float(arrow.group(1))


def test_same_short_code_on_two_islands_stays_two_substations(client):
    """KRSAN is Kraksaan in Jawa Timur (SS Paiton 1,2,3) and Keramasan in
    Sumsel; both 150 kV. Ingest used to reuse the first one it met, so the two
    became one node. A node held only by another system's subsystems is now a
    different site: Sumsel gets KRSAN@SUMATERA, drawn as plain KRSAN."""
    def publish(name, code):
        path = SAMPLE_XLSX.parent / name
        r = client.post("/api/ingest/parse-file", files={"file": (path.name, path.read_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 200, r.text
        draft = r.json()
        for n in draft["nodes"]:
            n["resolution"] = "NEW"
            n["canonical_id"] = None
        r = client.post('/api/ingest/publish', json={
            'draft': draft, 'subsystem_code': code, 'subsystem_name': code})
        assert r.status_code == 200, r.text

    publish('ss_paiton123_ingest.xlsx', 'SS_PAITON123')
    publish('ss_sumsel_ingest.xlsx', 'SS_SUMSEL')
    views = {v['view_key']: v['id'] for v in client.get('/api/views').json()}
    paiton = client.get(f"/api/views/{views['SS_PAITON123_FULL']}/graph").json()
    sumsel = client.get(f"/api/views/{views['SS_SUMSEL_FULL']}/graph").json()
    kraksaan = next(n for n in paiton['nodes'] if n.get('code') == 'KRSAN')
    keramasan = next(n for n in sumsel['nodes'] if n.get('code') == 'KRSAN@SUMATERA')
    assert kraksaan['id'] != keramasan['id']
    assert 'Keramasan' in keramasan['name'] and 'Kraksaan' in kraksaan['name']
    svg = client.get(f"/api/views/{views['SS_SUMSEL_FULL']}/sld.svg").text
    assert '>KRSAN<' in svg and 'KRSAN@SUMATERA<' not in svg


def test_prbc_three_views_preserve_multiple_continuations():
    from app.services.ingest_parser import parse_upload
    path = Path(__file__).parent / 'fixtures/jakban_legacy/ss_prbc_ingest.xlsx'
    data = parse_upload(path.read_bytes(), path.name)
    assert {v['view_key'] for v in data['subsystem']['views']} == {'BEKASI', 'PRIOK', 'CAWANG'}
    ancol = next(e for e in data['connections'] if {e['from_external_key'],e['to_external_key']} == {'ANGKE','ANCOL'})
    assert ancol['view_keys'] == ['PRIOK']
    plumpang = next(n for n in data['objects'] if n['external_key'] == 'PLPNG40')
    appearances = plumpang['bay_appearances']
    assert {(a['feeder_key'], tuple(a['view_keys'])) for a in appearances} >= {
        ('HNDAH', ('BEKASI',)), ('KDSPI', ('BEKASI',))}
    from scripts.audit_sample_workbooks import audit_one
    result = audit_one(path)
    assert result['ok'], result
    assert len(result['views']) == 3


def test_ibt_keeps_each_units_explicit_endpoints(monkeypatch):
    import openpyxl
    from app.services.ingest_parser import parse_xlsx

    class Sheet:
        def __init__(self, rows): self.rows = rows
        def iter_rows(self, values_only=True): return iter(self.rows)

    class Book(dict):
        @property
        def sheetnames(self): return list(self)

    book = Book({
        'Gardu_Induk_dan_Aset': Sheet([
            ('Kode', 'Tipe', 'Tegangan', 'Bus HV', 'Bus LV', 'No IBT', 'Sudut Pandang'),
            ('HV WITH SPACE', 'Busbar GITET', 500, None, None, None, 'A;B;C'),
            ('LV1', 'Busbar GI', 150, None, None, None, 'A'),
            ('LV2', 'Busbar GI', 70, None, None, None, 'B'),
            ('IBT 1', 'IBT n-Winding', None, 'HV WITH SPACE', 'LV1', 1, 'A'),
            ('IBT 2', 'IBT n-Winding', None, 'HV WITH SPACE', 'LV2', 2, 'B'),
            ('IBT 3', 'IBT n-Winding', None, 'HV WITH SPACE', 'MISSING', 3, 'C')]),
        'Jalur_Transmisi': Sheet([('Dari GI', 'Ke GI')]),
    })
    monkeypatch.setattr(openpyxl, 'load_workbook', lambda *a, **kw: book)
    with pytest.raises(ValueError, match='endpoint tak dikenal'):
        parse_xlsx(b'', 'ibt.xlsx')
    book['Gardu_Induk_dan_Aset'].rows.pop()
    parsed = parse_xlsx(b'', 'ibt.xlsx')
    links = parsed['connections']
    assert [(e['from_external_key'], e['to_external_key'], e['unit_no'], e['view_keys'])
            for e in links] == [
                ('HV WITH SPACE', 'LV1', '1', ['A']),
                ('HV WITH SPACE', 'LV2', '2', ['B'])]


@pytest.fixture()
def client():
    fd, path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    os.environ["DATABASE_URL"] = f"sqlite:///{path}"
    import app.db as db_mod
    import app.models
    importlib.reload(db_mod)
    for name in ("app.models", "app.services.topology", "app.services.reconciliation",
                 "app.services.sld_renderer", "app.services.ingestion",
                 "app.services.ingest_parser", "app.services.ingest",
                 "app.services.editor", "app.services.editor_risks",
                 "app.services.seed_ss_lbk", "app.services.seed_ss_bll",
                 "app.services.seed", "app.api.routes", "app.main"):
        importlib.reload(importlib.import_module(name))
    from fastapi.testclient import TestClient
    import app.main as main_mod
    with TestClient(main_mod.app) as c:
        yield c
    try:
        os.unlink(path)
    except OSError:
        pass


def _draft(client, fmt="xlsx"):
    if fmt == "xlsx":
        r = client.post("/api/ingest/parse-file",
                        files={"file": ("ss_cwd_ingest.xlsx", SAMPLE_XLSX.read_bytes(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    else:
        payload = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
        r = client.post("/api/ingest/parse", json={"payload": payload})
    assert r.status_code == 200, r.text
    return r.json()


def test_parse_xlsx_template_builds_draft(client):
    """The 3-sheet PLN template (Jalur_Transmisi + Gardu_Induk_dan_Aset +
    Data_Kerawanan_Detail) is the primary format."""
    d = _draft(client, "xlsx")
    assert d["subsystem"]["code"] == "SS_CWD"          # from the Info sheet
    assert len(d["nodes"]) == 15
    assert len(d["edges"]) == 13
    assert len(d["risks"]) == 4
    # GITET busbars are split from their 150 kV bus
    keys = {n["external_key"] for n in d["nodes"]}
    assert "GITET_DEPOK" in keys and "DEPOK" in keys
    # IBT links carry a unit number AND feed the right 150 kV bus:
    # GITET Cawang -> CWBRU (Cawang Baru), not CWANG (Cawang Lama, a Tier-2 GI)
    ibt = [e for e in d["edges"] if e["circuit_type_hint"] == "IBT_LINK"]
    assert len(ibt) == 3
    cwang_ibt = [e for e in ibt if e["from_key"] == "GITET_CWANG"]
    assert cwang_ibt and all(e["to_key"] == "CWBRU" for e in cwang_ibt)
    assert d["validation"]["ok"] is True


def test_zero_based_workbook_tiers_are_fully_converted(client):
    """Every value in `Tier (Mulai 0)` is zero-based, including tiers > 0."""
    d = _draft(client, "xlsx")
    tiers = {n["external_key"]: n["tier_hint"] for n in d["nodes"]}
    assert tiers["GITET_CWANG"] == 1
    assert tiers["CWBRU"] == 1
    assert tiers["CWANG"] == 2
    assert tiers["STBDI"] == 3


def test_lbk_template_preserves_single_phi_planned_ibt_and_view_scoped_boundary():
    from app.services.ingest_parser import parse_upload

    legacy = Path(__file__).parent / "fixtures/jakban_legacy/ss_lbk_ingest.xlsx"
    payload = parse_upload(legacy.read_bytes(), legacy.name)
    single = {(c["from_external_key"], c["to_external_key"])
              for c in payload["connections"] if c.get("single_phi")}
    assert {("SNYAN", "GISPD"), ("GISPD", "DNYSA"),
            ("PSKMS", "PSKBR"), ("PSKBR", "GJTGL"),
            ("GJTGL", "PSKMS")} <= single

    nckupa = next(c for c in payload["connections"]
                   if c["from_external_key"] == "NCKUPA" and c["relation_type"] == "IBT_LINK")
    assert nckupa["status_hint"] == "PLANNED"
    assert nckupa["view_keys"] == ["BALARAJA"]

    dksbi = next(o for o in payload["objects"] if o["external_key"] == "DKSBI")
    assert dksbi["role_hint"] == "BOUNDARY"
    assert dksbi["is_bay"] is False
    assert dksbi["bay_view_keys"] == ["KEMBANGAN"]


def test_ingest_multiple_ibt_links_do_not_collide(client):
    """A GITET with several IBT links -- and a hand-added extra one -- must
    materialise without a UNIQUE-constraint crash on subsystem_membership."""
    d = _draft(client, "xlsx")
    d["edges"].append({
        "from_key": "GITET_CWANG", "to_key": "CWBRU", "relation_type": "IBT_LINK",
        "circuit_type_hint": "IBT_LINK", "status_hint": "ENERGIZED",
        "circuit_count": 1, "unit_no": "4", "confidence": 1.0, "confirmed": True})
    r = client.post("/api/ingest/preview.svg", json=d)
    assert r.status_code == 200 and b"<svg" in r.content
    assert b"UNIQUE constraint" not in r.content

    pub = client.post("/api/ingest/publish", json={
        "draft": d, "subsystem_code": "SS_CWD", "subsystem_name": "Cawang 2,3 - Depok 1",
        "effective_date": "2026-06-30"})
    assert pub.status_code == 200, pub.text


def test_bay_stub_preserves_single_or_double_circuit_count(client):
    d = _draft(client, "xlsx")
    bay = next(n for n in d["nodes"] if n["external_key"] == "LKONG2")
    bay["bay_circuit_count"] = 2
    r = client.post("/api/ingest/preview.svg", json=d)
    assert r.status_code == 200, r.text
    root = ET.fromstring(r.text)
    ns = {"s": "http://www.w3.org/2000/svg"}
    stub = root.find('.//s:g[@class="sld-bay"][@data-code="LKONG2"]', ns)
    assert stub is not None and stub.get("data-circuit-count") == "2"
    assert len(stub.findall('s:path', ns)) == 2


def test_parse_json_handoff_still_works(client):
    d = _draft(client, "json")
    assert len(d["nodes"]) == 15 and len(d["edges"]) == 13
    dc = next(e for e in d["edges"] if e["from_key"] == "DEPOK" and e["to_key"] == "CWANG")
    assert dc["needs_review"] is False
    gd = next(n for n in d["nodes"] if n["external_key"] == "GITET_DEPOK")
    assert gd["resolution"] == "NEW"


def test_bali_sklt_is_two_separated_pairs_from_source_boundary(client):
    r = client.post("/api/ingest/parse-file", files={
        "file": (BALI_JSON.name, BALI_JSON.read_bytes(), "application/json")})
    assert r.status_code == 200, r.text
    d = r.json()
    bw = next(n for n in d["nodes"] if n["external_key"] == "BANYUWANGI")
    assert bw["tier_hint"] is None and bw["role_hint"] == "SOURCE_BOUNDARY"
    sklt = [e for e in d["edges"]
            if {e["from_key"], e["to_key"]} == {"BANYUWANGI", "GILIMANUK"}]
    assert [(e["unit_no"], e["circuit_count"]) for e in sklt] == [("1,2", 2), ("3,4", 2)]
    ais_gis = next(e for e in d["edges"]
                   if {e["from_key"], e["to_key"]} == {"PESANGGARAN", "GIS_PESANGGARAN"})
    assert ais_gis["circuit_type_hint"] == "SKTT" and ais_gis["circuit_count"] == 2
    risk = next(x for x in d["risks"] if x["seq_no"] == 3)
    assert risk["pin_key"] == "BANYUWANGI-GILIMANUK:1,2"

    preview = client.post("/api/ingest/preview.svg", json=d)
    assert preview.status_code == 200, preview.text
    root = ET.fromstring(preview.text)
    ns = {"s": "http://www.w3.org/2000/svg"}
    assert root.find('.//s:g[@id="busbars"]/s:g[@data-code="BANYUWANGI"]', ns) is None
    boundary = root.find('.//s:g[@id="bays"]/s:g[@data-code="BANYUWANGI"]', ns)
    assert boundary is not None and boundary.get("data-circuit-count") == "4"
    stub_paths = boundary.findall('s:path', ns)
    assert len(stub_paths) == 4
    assert all(float(__import__('re').findall(r'-?\d+(?:\.\d+)?', p.get('d'))[-1]) < 210
               for p in stub_paths)
    for group in root.findall('.//s:g[@id="circuits"]/s:g', ns):
        if "BANYUWANGI_GILIMANUK" in (group.get("data-circuit-code") or ""):
            continue  # the one permitted source lead through the Tier-0 strip
        for path in group.findall('s:path[@class="sld-wire"]', ns):
            xy = list(map(float, __import__('re').findall(r'-?\d+(?:\.\d+)?', path.get('d'))))
            points = list(zip(xy[::2], xy[1::2]))
            assert not any(a[1] == b[1] and a[1] < 210 for a, b in zip(points, points[1:]))


def test_ingest_flags_an_existing_subsystem(client):
    """Uploading the datasheet of a subsystem that already exists (SS_BLL) must
    be blocked -- the user should edit it via /editor, not bootstrap a copy."""
    from app.services.ingest_parser import normalise
    from app.services import ingest as ingest_svc
    import app.db as db_mod

    # a tiny payload made of GIs that are already SS_BLL members
    payload = normalise({
        "subsystem": {"code": "SS_WHATEVER", "name": "x"},
        "objects": [
            {"external_key": "NBRJA", "object_type": "GI", "raw_label": "NBRJA", "tier_hint": 1},
            {"external_key": "LSTEL", "object_type": "GI", "raw_label": "LSTEL", "tier_hint": 2},
            {"external_key": "SPMIL", "object_type": "GI", "raw_label": "SPMIL", "tier_hint": 3},
        ],
        "connections": [
            {"from_external_key": "NBRJA", "to_external_key": "LSTEL"},
            {"from_external_key": "LSTEL", "to_external_key": "SPMIL"},
        ],
        "risks": [],
    })
    with db_mod.SessionLocal() as db:
        d = ingest_svc.build_draft(db, payload)
    assert d["existing"] is not None
    assert d["existing"]["subsystem_code"] == "SS_BLL"
    assert d["validation"]["ok"] is False
    assert any("SS_BLL" in p and "/editor" in p for p in d["validation"]["problems"])


def test_parse_writes_nothing_to_db(client):
    _draft(client)
    assert not any(s["code"] == "SS_CWD" for s in client.get("/api/subsystems").json())
    assert not any(v["view_key"] == "SS_CWD_FULL" for v in client.get("/api/views").json())


def test_coordinates_survive_draft_publish_and_graph(client):
    payload = {
        "subsystem": {"code": "SS_GEO", "name": "Geo fixture"},
        "objects": [
            {"external_key": "GEOW", "object_type": "GI", "raw_label": "Geo West",
             "tier_hint": 1, "latitude": -6.2, "longitude": 106.5},
            {"external_key": "GEOE", "object_type": "GI", "raw_label": "Geo East",
             "tier_hint": 2, "latitude": -6.3, "longitude": 107.5},
        ],
        "connections": [{"from_external_key": "GEOW", "to_external_key": "GEOE"}],
        "risks": [],
    }
    draft = client.post("/api/ingest/parse", json=payload).json()
    assert draft["nodes"][0]["latitude"] == -6.2
    published = client.post("/api/ingest/publish", json={
        "draft": draft, "subsystem_code": "SS_GEO", "subsystem_name": "Geo fixture",
    })
    assert published.status_code == 200, published.text
    graph = client.get(f"/api/views/{published.json()['view_id']}/graph").json()
    coords = {n["code"]: (n["latitude"], n["longitude"]) for n in graph["nodes"]
              if n["kind"] == "SUBSTATION"}
    assert coords == {"GEOW": (-6.2, 106.5), "GEOE": (-6.3, 107.5)}


def test_preview_renders_from_blob_and_rolls_back(client):
    d = _draft(client)
    r = client.post("/api/ingest/preview.svg", json=d)
    assert r.status_code == 200
    assert r.headers["content-type"] == "image/svg+xml"
    assert b"<svg" in r.content
    # canon still clean after a preview
    assert not any(s["code"] == "SS_CWD" for s in client.get("/api/subsystems").json())


def test_preview_never_500s_on_a_broken_draft(client):
    d = _draft(client)
    # duplicate a code -> materialise would hit the UNIQUE constraint
    d["nodes"][3]["confirmed_code"] = d["nodes"][2]["confirmed_code"]
    r = client.post("/api/ingest/preview.svg", json=d)
    assert r.status_code == 200
    assert b"<svg" in r.content
    assert "dipakai dua node".encode() in r.content


def test_validate_blocks_unresolved_node(client):
    d = _draft(client)
    for n in d["nodes"]:
        if n["external_key"] == "TRSNA":
            n["resolution"] = "NEW"
            n["confirmed_code"] = ""
    v = client.post("/api/ingest/validate", json=d).json()
    assert v["ok"] is False
    assert any("belum diberi kode" in p for p in v["problems"])


def test_publish_materialises_subsystem(client):
    d = _draft(client)
    pub = client.post("/api/ingest/publish", json={
        "draft": d, "subsystem_code": "SS_CWD",
        "subsystem_name": "Cawang 2,3 - Depok 1", "effective_date": "2026-06-30"})
    assert pub.status_code == 200, pub.text
    assert pub.json()["view_key"] == "SS_CWD_FULL"

    views = client.get("/api/views").json()
    cwd = next(v for v in views if v["view_key"] == "SS_CWD_FULL")
    g = client.get(f"/api/views/{cwd['id']}/graph").json()
    subs = {n["code"]: n for n in g["nodes"] if n["kind"] == "SUBSTATION"}
    assert len(subs) == 15
    # Tier bands match Buku Kerawanan Sec 2.7
    assert subs["DEPOK"]["tier"] == 1 and subs["CWBRU"]["tier"] == 1
    assert subs["CWANG"]["tier"] == 2 and subs["ABDGP"]["tier"] == 2
    assert subs["STBDI"]["tier"] == 3 and subs["DNYSA"]["tier"] == 3
    assert len(g["overlays"]["risk"]) == 4
    assert all(r["attach_kind"] for r in g["overlays"]["risk"])
    assert client.get(f"/api/views/{cwd['id']}/sld.svg").status_code == 200

    # a traced (0.5) ruas carries a NEEDS_REVIEW note on the canonical circuit
    import app.db as db_mod
    from app.models import Circuit
    with db_mod.SessionLocal() as db:
        c = db.query(Circuit).filter(Circuit.code == "SUTT_CWBRU_DRNTG").first()
        assert c is not None and (c.note or "").startswith("NEEDS_REVIEW")
        assert c.confidence < 0.9


def test_publish_refuses_duplicate_subsystem(client):
    d = _draft(client)
    body = {"draft": d, "subsystem_code": "SS_CWD",
            "subsystem_name": "x", "effective_date": None}
    assert client.post("/api/ingest/publish", json=body).status_code == 200
    r2 = client.post("/api/ingest/publish", json=body)
    assert r2.status_code == 400
    assert "sudah ada" in r2.json()["detail"]


def test_same_site_code_at_different_voltages_materialises_as_distinct_buses(client):
    def publish(code, kv, peer):
        payload = {
            "subsystem": {"code": code, "name": code},
            "objects": [
                {"external_key": "CURUG", "object_type": "GI", "raw_label": "Curug",
                 "site_name": "Curug", "voltage_hv_kv": kv, "tier_hint": 1},
                {"external_key": peer, "object_type": "GI", "raw_label": peer,
                 "voltage_hv_kv": kv, "tier_hint": 2},
            ],
            "connections": [{"from_external_key": "CURUG", "to_external_key": peer,
                             "voltage_hv_kv": kv}],
            "risks": [],
        }
        draft = client.post("/api/ingest/parse", json=payload).json()
        response = client.post("/api/ingest/publish", json={
            "draft": draft, "subsystem_code": code, "subsystem_name": code})
        assert response.status_code == 200, response.text
        return response.json()["view_id"]

    publish("SS_CURUG150", 150, "PEER150")
    view70 = publish("SS_CURUG70", 70, "PEER70")
    graph = client.get(f"/api/views/{view70}/graph").json()
    curug = next(n for n in graph["nodes"] if n["kind"] == "SUBSTATION"
                 and n["name"] == "Curug")
    assert curug["code"] == "CURUG_70KV"
    assert curug["voltage_kv"] == 70


def test_publish_refuses_invalid_draft(client):
    d = _draft(client)
    for n in d["nodes"]:
        if n["external_key"] == "MPLMA":
            n["confirmed_code"] = ""
    r = client.post("/api/ingest/publish", json={
        "draft": d, "subsystem_code": "SS_CWD", "subsystem_name": "x",
        "effective_date": None})
    assert r.status_code == 400
    assert "validasi gagal" in r.json()["detail"]


def test_parser_rejects_payload_without_subsystem(client):
    r = client.post("/api/ingest/parse", json={"payload": {"objects": []}})
    assert r.status_code == 400
    assert "wajib" in r.json()["detail"]


def test_explicit_symbol_counts_survive_excel_draft_preview_publish(client, monkeypatch):
    import openpyxl
    from app.services.ingest_parser import parse_xlsx

    class Sheet:
        def __init__(self, rows):
            self.rows = rows
        def iter_rows(self, values_only=True):
            return iter(self.rows)

    class Book(dict):
        @property
        def sheetnames(self):
            return list(self)

    book = Book({
        'Info': Sheet([('Kode Subsistem', 'SS_SYMBOLS'), ('Nama Subsistem', 'Symbol test')]),
        'Gardu_Induk_dan_Aset': Sheet([
            ('Kode', 'Tipe', 'Tier', 'Jumlah Trafo', 'Jumlah Kapasitor'),
            ('SYMSRC', 'Busbar GI', 1, 0, 0), ('SYMLOAD', 'Busbar GI', 2, 3, 2)]),
        'Jalur_Transmisi': Sheet([('Dari GI', 'Ke GI', 'Jumlah Sirkit'), ('SYMSRC', 'SYMLOAD', 2)]),
    })
    monkeypatch.setattr(openpyxl, 'load_workbook', lambda *a, **kw: book)
    payload = parse_xlsx(b'', 'symbols.xlsx')
    draft = client.post('/api/ingest/parse', json={'payload': payload}).json()
    load = next(n for n in draft['nodes'] if n['external_key'] == 'SYMLOAD')
    assert (load['transformer_count'], load['capacitor_count']) == (3, 2)
    response = client.post('/api/ingest/preview.svg', json=draft)
    assert response.status_code == 200
    assert response.text.count('class="load-transformer"') == 3
    assert response.text.count('class="shunt-capacitor"') == 2
    published = client.post('/api/ingest/publish', json={
        'draft': draft, 'subsystem_code': 'SS_SYMBOLS', 'subsystem_name': 'Symbol test'})
    assert published.status_code == 200, published.text
    svg = client.get(f'/api/views/{published.json()["view_id"]}/sld.svg').text
    assert svg.count('class="load-transformer"') == 3
    assert svg.count('class="shunt-capacitor"') == 2


def test_negative_symbol_count_is_not_published(client):
    draft = _draft(client, 'json')
    draft['nodes'][0]['capacitor_count'] = -1
    result = client.post('/api/ingest/validate', json=draft).json()
    assert not result['ok']
    assert any('jumlah simbol' in p for p in result['problems'])


def test_line_voltage_reaches_read_only_audit(monkeypatch):
    import openpyxl
    from app.services.ingest_parser import parse_xlsx

    class Sheet:
        def __init__(self, rows): self.rows = rows
        def iter_rows(self, values_only=True): return iter(self.rows)
    class Book(dict):
        @property
        def sheetnames(self): return list(self)

    book = Book({
        'Info': Sheet([('Kode Subsistem', 'SS_VOLT'), ('Nama Subsistem', 'Voltage audit')]),
        'Gardu_Induk_dan_Aset': Sheet([
            ('Kode', 'Tipe', 'Tier', 'Tegangan'),
            ('A', 'Busbar GI', 1, '150 kV'), ('B', 'Busbar GI', 2, '150 kV')]),
        'Jalur_Transmisi': Sheet([
            ('Dari GI', 'Ke GI', 'Tegangan', 'Jumlah Sirkit'),
            ('A', 'B', '20 kV', 1)]),
    })
    monkeypatch.setattr(openpyxl, 'load_workbook', lambda *a, **kw: book)
    payload = parse_xlsx(b'', 'voltage.xlsx')
    assert payload['connections'][0]['voltage_hv_kv'] == 20
    from pathlib import Path
    from scripts.audit_voltage_consistency import audit_payloads
    findings = audit_payloads([(Path('voltage.xlsx'), payload)])
    assert any(f['rule'] == 'LINE_BUS_VOLTAGE_MISMATCH' for f in findings)


def test_voltage_audit_finds_cross_workbook_collision_and_orphan_generator():
    from pathlib import Path
    from scripts.audit_voltage_consistency import audit_payloads

    def payload(code, kv, generator=False):
        objects = [{
            'external_key': 'CURUG', 'object_type': 'GI', 'raw_label': 'Curug',
            'site_name': 'Curug', 'voltage_hv_kv': kv,
        }]
        if generator:
            objects.append({
                'external_key': 'KIT_X', 'object_type': 'GENERATING_UNIT',
                'raw_label': 'PLTA X', 'site_name': 'PLTA X',
                'voltage_hv_kv': 11, 'outlet_key': None,
            })
        return {'subsystem': {'code': code}, 'objects': objects, 'connections': []}

    findings = audit_payloads([
        (Path('a.xlsx'), payload('SS_A', 150, generator=True)),
        (Path('b.xlsx'), payload('SS_B', 70)),
    ])
    rules = {f['rule'] for f in findings}
    assert 'CROSS_WORKBOOK_CODE_VOLTAGE_COLLISION' in rules
    assert 'GENERATOR_WITHOUT_OUTLET' in rules


def test_multiview_500kv_and_multi_risk_tags_publish(client, monkeypatch):
    import openpyxl
    from app.services.ingest_parser import parse_xlsx

    class Sheet:
        def __init__(self, rows): self.rows = rows
        def iter_rows(self, values_only=True): return iter(self.rows)
    class Book(dict):
        @property
        def sheetnames(self): return list(self)

    book = Book({
        'Info': Sheet([('Kode Subsistem', 'SS_500_MULTI'), ('Nama Subsistem', '500 kV multiview')]),
        'Views': Sheet([('Kode View', 'Nama View', 'Sumber Tier-1 (kode GI, pisah ;)'),
                        ('V1', 'View one', 'A'), ('V2', 'View two', 'B'),
                        ('V3', 'Shared source, separate branch', 'A')]),
        'Gardu_Induk_dan_Aset': Sheet([
            ('Kode', 'Tipe', 'Tier', 'Tegangan', 'No Kerawanan', 'Sudut Pandang', 'Bus Terhubung'),
            ('A', 'Busbar GITET', 1, '500 kV', None, 'V1', None),
            ('X', 'Busbar GITET', 2, '500 kV', '5;7', 'V1', None),
            ('GEN', 'Pembangkit', 1, '500 kV', None, 'V1', 'A'),
            ('B', 'Busbar GITET', 1, '500 kV', None, 'V2', None),
            ('Y', 'Busbar GITET', 2, '500 kV', None, 'V2', None),
            ('Z', 'Busbar GITET', 2, '500 kV', None, 'V3', None)]),
        'Jalur_Transmisi': Sheet([
            ('Dari GI', 'Ke GI', 'Jumlah Sirkit', 'No Kerawanan', 'Sudut Pandang'),
            ('A', 'X', 2, 7, 'V1'), ('B', 'Y', 2, None, 'V2'),
            ('A', 'Z', 2, None, 'V3')]),
        'Data_Kerawanan_Detail': Sheet([
            ('No', 'Kondisi / Permasalahan'), (5, 'Risk at X'), (7, 'Risk on A-X')]),
    })
    monkeypatch.setattr(openpyxl, 'load_workbook', lambda *a, **kw: book)
    payload = parse_xlsx(b'', 'multi.xlsx')
    assert [v['view_key'] for v in payload['subsystem']['views']] == ['V1', 'V2', 'V3']
    assert next(o for o in payload['objects'] if o['external_key'] == 'GEN')['object_type'] == 'GENERATING_UNIT'
    assert [(r['seq_no'], r['pin_kind'], r['pin_key']) for r in payload['risks']] == [
        (5, 'SUBSTATION', 'X'), (7, 'CIRCUIT', 'A-X')]
    draft = client.post('/api/ingest/parse', json={'payload': payload}).json()
    assert draft['validation']['ok'], draft['validation']['problems']
    published = client.post('/api/ingest/publish', json={
        'draft': draft, 'subsystem_code': 'SS_500_MULTI', 'subsystem_name': '500 kV multiview'})
    assert published.status_code == 200, published.text
    assert {v['view_key'] for v in published.json()['views']} == {'SS_500_MULTI_V1', 'SS_500_MULTI_V2', 'SS_500_MULTI_V3'}
    views = {v['view_key']: v for v in client.get('/api/views').json()}
    g1 = client.get(f"/api/views/{views['SS_500_MULTI_V1']['id']}/graph").json()
    g2 = client.get(f"/api/views/{views['SS_500_MULTI_V2']['id']}/graph").json()
    assert {n['code'] for n in g1['nodes']} >= {'A', 'X', 'GEN'}
    assert {n['code'] for n in g2['nodes']} >= {'B', 'Y'}
    assert all(n['code'] not in {'B', 'Y'} for n in g1['nodes'])
    g3 = client.get(f"/api/views/{views['SS_500_MULTI_V3']['id']}/graph").json()
    assert {n['code'] for n in g3['nodes']} == {'A', 'Z'}
    assert next(n['id'] for n in g1['nodes'] if n['code'] == 'A') == next(
        n['id'] for n in g3['nodes'] if n['code'] == 'A')


def test_one_risk_number_on_several_objects_pins_each_of_them(client, monkeypatch):
    """The book draws one numbered starburst on every object a finding sits on.
    In a sheet that opts in with "Multi Pin = Ya", a number written on several
    objects keeps them all: the last one written stays the primary (what a
    single-pin risk always resolved to), the rest become extra pins, and the
    SLD marks each one. Without the flag the sheet keeps its single pin."""
    import openpyxl
    from app.services.ingest_parser import parse_xlsx

    class Sheet:
        def __init__(self, rows): self.rows = rows
        def iter_rows(self, values_only=True): return iter(self.rows)
    class Book(dict):
        @property
        def sheetnames(self): return list(self)

    info = [('Kode Subsistem', 'SS_MULTIPIN'), ('Nama Subsistem', 'Multi pin')]
    book = Book({
        'Info': Sheet(info),
        'Gardu_Induk_dan_Aset': Sheet([
            ('Kode', 'Tipe', 'Tier', 'Tegangan', 'Bus HV', 'Bus LV', 'No IBT', 'No Kerawanan'),
            ('G', 'Busbar GITET', 1, '500 kV', None, None, None, None),
            ('IBT 1 G', 'IBT 3-Winding', 2, '500/150 kV', 'G', 'A', 1, 3),
            ('IBT 2 G', 'IBT 3-Winding', 2, '500/150 kV', 'G', 'A', 2, 3),
            ('A', 'Busbar GI', 1, '150 kV', None, None, None, None),
            ('B', 'Busbar GI', 2, '150 kV', None, None, None, 2),
            ('C', 'Busbar GI', 3, '150 kV', None, None, None, 1)]),
        'Jalur_Transmisi': Sheet([
            ('Dari GI', 'Ke GI', 'Jumlah Sirkit', 'No Kerawanan'),
            ('A', 'B', 2, 1), ('B', 'C', 1, 1)]),
        'Data_Kerawanan_Detail': Sheet([
            ('No', 'Kondisi / Permasalahan'),
            (1, 'Radial A - B - C'), (2, 'Only at B'), (3, 'IBT 1,2 G')]),
    })
    monkeypatch.setattr(openpyxl, 'load_workbook', lambda *a, **kw: book)
    # without the flag: same primary, no extra pins (the Jamali sheets)
    single = {r['seq_no']: r for r in parse_xlsx(b'', 'multipin.xlsx')['risks']}
    assert (single[1]['pin_kind'], single[1]['pin_key']) == ('CIRCUIT', 'B-C')
    assert all(r['extra_pins'] == [] for r in single.values())

    info.append(('Multi Pin', 'Ya'))
    payload = parse_xlsx(b'', 'multipin.xlsx')
    by_seq = {r['seq_no']: r for r in payload['risks']}
    assert (by_seq[1]['pin_kind'], by_seq[1]['pin_key']) == ('CIRCUIT', 'B-C')
    assert by_seq[1]['extra_pins'] == [['SUBSTATION', 'C'], ['CIRCUIT', 'A-B']]
    assert by_seq[2]['extra_pins'] == []
    assert (by_seq[3]['pin_kind'], by_seq[3]['pin_key']) == ('TRANSFORMER', 'G:2')
    assert by_seq[3]['extra_pins'] == [['TRANSFORMER', 'G:1']]

    draft = client.post('/api/ingest/parse', json={'payload': payload}).json()
    assert draft['validation']['ok'], draft['validation']['problems']
    published = client.post('/api/ingest/publish', json={
        'draft': draft, 'subsystem_code': 'SS_MULTIPIN', 'subsystem_name': 'Multi pin'})
    assert published.status_code == 200, published.text

    view = next(v for v in client.get('/api/views').json() if v['view_key'].startswith('SS_MULTIPIN'))
    graph = client.get(f"/api/views/{view['id']}/graph").json()
    risks = {r['seq_no']: r for r in graph['overlays']['risk']}
    assert sorted(a['attach_label'] for a in risks[1]['attachments']) == ['A-B', 'C']
    assert [a['attach_label'] for a in risks[3]['attachments']] == ['G:1']

    svg = ET.fromstring(client.get(f"/api/views/{view['id']}/sld.svg").text)
    pins = [g.get('data-risk-seqs') for g in svg.iter('{http://www.w3.org/2000/svg}g')
            if g.get('class') == 'risk-pin']
    # risk 1 on C, A-B and B-C; IBT 1 and 2 share one spot, so one pin for #3
    assert sum('1' in (s or '').split(',') for s in pins) == 3
    assert sum('3' in (s or '').split(',') for s in pins) == 1


def test_an_ibt_bay_on_the_500kv_map_draws_as_a_transformer_with_its_pin(client):
    """The Sistem 500 kV IBT map hangs each GITET's IBTs off it as Bay rows
    (Jenis = IBT). The parser used to drop `Jenis`, so they drew as plain
    feeder stubs, and a finding sited at one had no pin at all. Now each is a
    transformer whose secondary is cut, labelled with its units, and pinned."""
    import re
    path = SAMPLE_XLSX.parent / 'system_ibt_500_ingest.xlsx'
    r = client.post("/api/ingest/parse-file", files={"file": (path.name, path.read_bytes(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    draft = r.json()
    assert any(n.get("bay_kind") == "IBT" for n in draft["nodes"])
    for n in draft["nodes"]:
        n["resolution"] = "NEW"
    published = client.post('/api/ingest/publish', json={
        'draft': draft, 'subsystem_code': 'SYSTEM_IBT_500', 'subsystem_name': 'IBT 500/150 kV'})
    assert published.status_code == 200, published.text
    view = next(v for v in client.get('/api/views').json() if v['view_key'].startswith('SYSTEM_IBT_500'))
    svg = client.get(f"/api/views/{view['id']}/sld.svg").text
    from tests.test_sld_geometry import geometry_errors
    assert geometry_errors(svg) == []
    bay = re.search(r'<g class="sld-bay sld-bay-ibt"[^>]*data-code="IBT_TMBUN7".*?</text>', svg, re.S)
    assert bay, "Tambun's IBT is drawn as an IBT bay"
    assert bay.group(0).count('<circle') >= 3            # three windings, not a stub's dot
    assert '>IBT 1,2</text>' in bay.group(0)
    assert re.search(r'class="risk-pin" data-risk-seqs="16"', svg)   # #16 is sited on it
