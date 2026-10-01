"""Read-only, source-grounded audit of the entire Krian 1,2-Gresik 1,2 SS.

Run: python -X utf8 scripts/audit_krian12_gresik.py
This does not certify live operating state or modify either workbook.
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / 'samples/ss_krian12_gresik_ingest.xlsx'
EXPORT = ROOT.parent / 'template_kerawanan_subsistem_krian12_gresik12.xlsx'
TEMPLATE = ROOT.parent / 'template_kerawanan_subsistem_suralaya_cilegon (1).xlsx'
OUT = ROOT / 'output/audit-krian12-gresik12'

# Independently read from Tabel 5.1, not from the builder's pin assignments.
# #13 deliberately remains unresolved: table says Waru-GIS Buduran, while
# the figure draws a pin below BDRAN5 and the export pins three segments.
EXPECTED_PINS = {
    1: {'IBT 1 KRIAN7', 'IBT 2 KRIAN7'},
    2: {'IBT 1 GRSIK7', 'IBT 2 GRSIK7'},
    3: {'SBRAT5--KLANG'},
    4: {'SBRAT5--SWHAN', 'SWHAN--GNSRI', 'GNSRI--WARU5'},
    5: {'SBRAT5--DARMO', 'DARMO--WARU5'},
    6: {'GRSIK5--TNDES'}, 7: {'GRLMA--SKREP'},
    8: {'TNDES--SWHAN'}, 9: {'SWHAN'},
    10: {'KLANG--RNKUT'}, 11: {'SKREP--WARU5'},
    12: {'WARU5--RNKUT'}, 14: {'BDRAN5'}, 15: {'MPION'},
    16: {'RNKUT--SLILO'}, 17: {'SLILO--KJRAN'},
    18: {'TNDES--PERAK', 'PERAK--UJUNG', 'TNDES--UJUNG'},
    19: {'PERAK'}, 20: {'KJRAN--GLMUR', 'GLMUR--BKLAN'},
    21: {'BKLAN'},
    22: {'UJUNG--KDING', 'UJUNG--BKLAN', 'KDING--BKLAN'},
    23: {'BKLAN--SAMPG'}, 24: {'SAMPG--PKSAN'},
    25: {'KJRAN', 'GULUK'},
}
SINGLE_PHI = {'TNDES--PERAK', 'PERAK--UJUNG', 'KJRAN--GLMUR', 'GLMUR--BKLAN'}
UNKNOWN_NAMES = {'ALTAP', 'KPANG', 'PTISM', 'GBONG', 'HJAYA', 'JSTEL', 'KSARI'}


def rows(ws):
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)
            if any(v is not None for v in r)]


def nums(value):
    return {int(n) for n in re.findall(r'\d+', str(value or ''))}


def edge(row):
    return f"{row['Dari GI']}--{row['Ke GI']}"


def same_edge(a, b):
    return frozenset(a.split('--')) == frozenset(b.split('--'))


def cell(value):
    return str(value or '-').replace('|', '\\|').replace('\n', ' ')


def audit():
    wb = openpyxl.load_workbook(SAMPLE, data_only=True)
    assets = rows(wb['Gardu_Induk_dan_Aset'])
    lines = rows(wb['Jalur_Transmisi'])
    risks = rows(wb['Data_Kerawanan_Detail'])
    codes = {a['Kode Singkatan'] for a in assets}
    pins = defaultdict(set)
    for a in assets:
        for n in nums(a['No Kerawanan']):
            pins[n].add(a['Kode Singkatan'])
    for ln in lines:
        for n in nums(ln['No Kerawanan']):
            pins[n].add(edge(ln))
    missing_ends = [edge(ln) for ln in lines
                    if ln['Dari GI'] not in codes or ln['Ke GI'] not in codes]
    missing_ibt_ends = [a['Kode Singkatan'] for a in assets
                        if str(a['Tipe Asset']).startswith('IBT')
                        and (a['Bus HV'] not in codes or a['Bus LV'] not in codes)]
    risk_checks = []
    for r in risks:
        n = int(r['No'])
        expected = EXPECTED_PINS.get(n)
        actual = pins[n]
        status = 'COCOK LOKASI TABEL'
        note = 'Kecocokan lokasi kerawanan; bukan validasi konfigurasi operasi.'
        if n == 13:
            status = 'PERLU REVIEW'
            note = ('Tabel: Waru-Buduran dan Waru-GIS Buduran. Gambar: pin di atas '
                    'dan di bawah BDRAN5. Alias GIS Buduran/NBRAN5 dan pin '
                    'WARU5-SDRJO/BDRAN5-NBRAN5 belum terkonfirmasi.')
        elif actual != expected:
            status = 'KONFLIK SUMBER'
            note = ('Gambar memiliki pin #23 juga di Kenjeran-Kedinding; tabel '
                    'hanya Bangkalan-Sampang. Jangan menggandakan temuan overload '
                    'ke Kenjeran-Kedinding tanpa klarifikasi pemilik dokumen.')
        if n in (9, 21):
            status = 'LOKASI COCOK; TOPOLOGI BELUM VALID'
            note = 'Kopel dibuka menurut tabel, tetapi model hanya satu node GI.'
        phi_edges = {'TNDES--PERAK', 'PERAK--UJUNG'} if n == 19 else {'KJRAN--GLMUR', 'GLMUR--BKLAN'}
        if n in (19, 20) and any(
            any(same_edge(edge(ln), e) for e in phi_edges)
            and str(ln['Single Phi']).lower() not in ('ya', 'yes', 'true', '1')
            for ln in lines
        ):
            status = 'LOKASI COCOK; SINGLE PHI HILANG'
            note = 'Tabel menyebut single phi; kolom Single Phi ruas terkait masih Tidak.'
        risk_checks.append(dict(no=n, actual=sorted(actual),
                                expected=sorted(expected) if expected else None,
                                status=status, note=note))

    ex = openpyxl.load_workbook(EXPORT, data_only=True)
    ea, el, er = rows(ex['Gardu_Induk']), rows(ex['Jalur_Transmisi']), rows(ex['Tabel_Kerawanan'])
    exported_counts = Counter(frozenset((ln['Dari GI'], ln['Ke GI'])) for ln in el)
    expected_counts = Counter()
    for ln in lines:
        expected_counts[frozenset((ln['Dari GI'], ln['Ke GI']))] += int(ln['Jumlah Sirkit'] or 1)
    for a in assets:
        if str(a['Tipe Asset']).startswith('IBT'):
            expected_counts[frozenset((a['Bus HV'], a['Bus LV']))] += 1
    template = openpyxl.load_workbook(TEMPLATE, data_only=True)
    headers_same = all([c.value for c in ex[s][1]] == [c.value for c in template[s][1]]
                       for s in template.sheetnames)
    guide_same = list(ex['Panduan_Simbol_SLD'].values) == list(template['Panduan_Simbol_SLD'].values)
    export_codes = {a['Kode Singkatan'] for a in ea}
    export_bad_ends = [edge(ln) for ln in el
                       if ln['Dari GI'] not in export_codes or ln['Ke GI'] not in export_codes]
    expected_line_pins = Counter()
    for ln in lines:
        for n in nums(ln['No Kerawanan']):
            expected_line_pins[(frozenset((ln['Dari GI'], ln['Ke GI'])), n)] += int(ln['Jumlah Sirkit'] or 1)
    for a in assets:
        if str(a['Tipe Asset']).startswith('IBT'):
            for n in nums(a['No Kerawanan']):
                expected_line_pins[(frozenset((a['Bus HV'], a['Bus LV'])), n)] += 1
    actual_line_pins = Counter((frozenset((ln['Dari GI'], ln['Ke GI'])), n)
                               for ln in el for n in nums(ln['No Kerawanan']))
    source_asset_pins = {a['Kode Singkatan']: nums(a['No Kerawanan']) for a in assets}
    export_asset_pins = {a['Kode Singkatan']: nums(a['No Kerawanan']) for a in ea}
    phi_missing = [edge(ln) for ln in lines
                   if any(same_edge(edge(ln), e) for e in SINGLE_PHI)
                   and str(ln['Single Phi']).lower() not in ('ya', 'yes', 'true', '1')]
    findings = [
        dict(id='F01', severity='BLOCKER', subject='Terminal Sawahan belum terpetakan',
             detail='Model/parser kini menyimpan seksi dan kopel OPEN untuk RNKUT, BKLAN, dan SWHAN. Terminal RNKUT/BKLAN dipetakan dari sumber; lima terminal SWHAN belum terpetakan sesuai konfirmasi pengguna. Analisis konektivitas subsistem tetap incomplete; lihat docs/bus-sections.md.'),
        dict(id='F02', severity='BLOCKER', subject='Pin #23',
             detail='Kenjeran-Kedinding membawa risiko overload #23 walau Tabel 5.1 hanya menyebut Bangkalan-Sampang. Konflik gambar/tabel harus dicatat, bukan dianggap dua fakta overload.'),
        dict(id='F03', severity='REVIEW', subject='Pin #13 / GIS Buduran',
             detail='Tiga ruas dipin: WARU5-SDRJO, WARU5-BDRAN5, BDRAN5-NBRAN5. Tabel hanya menamai Waru-Buduran dan Waru-GIS Buduran. Pastikan alias dan terminal/bay sebelum menetapkan pin.'),
        dict(id='F04', severity='ERROR', subject='Single Phi', detail=', '.join(phi_missing)),
        dict(id='F05', severity='REVIEW', subject='Garis putus-putus',
             detail='SKTT dan Beroperasi ditetapkan dari rupa garis dan dugaan tidak ada suplai lain. Gambar tanpa legenda jenis saluran/status tidak membuktikan kedua atribut itu. #3 SKTT eksplisit di tabel; ruas putus-putus lain perlu asset register/legenda.'),
        dict(id='F06', severity='REVIEW', subject='Identitas GI',
             detail='Nama resmi belum tersedia untuk ' + ', '.join(sorted(UNKNOWN_NAMES)) + '. Placeholder menjaga kode, tetapi belum merupakan master data terverifikasi. GITET vs bus 150 kV serta alias GIS Buduran perlu FunctLoc asli.'),
        dict(id='F07', severity='REVIEW', subject='Buduran 150/70 kV',
             detail='Gambar menampilkan unit 1 dan 7. Export mengisi No Sirkit 1 dan 2; unit 7 hanya tersimpan dalam nama baris/No IBT aset. Pastikan aplikasi menggunakan No IBT untuk identitas dan mendukung trafo 150/70 kV; legenda IBT template hanya menjelaskan 500/150 atau 275/150.'),
        dict(id='F08', severity='REVIEW', subject='Tingkat risiko dan status operasi',
             detail='Exporter memberi semua pin Sangat Rawan dan tanpa pin Normal. Itu aturan konversi, bukan penilaian severity dari buku. Seluruh ruas Beroperasi juga bukan telemetri atau konfirmasi keadaan 1 Oktober 2026.'),
        dict(id='F09', severity='LIMITATION', subject='Aplikasi teman / koordinat',
             detail='Kode ingest aplikasi teman tidak tersedia. Screenshot menampilkan Suralaya-Cilegon dan judul SUBSISTEM BOGOR; bukan bukti render Krian. Koordinat/map tidak diaudit dan tidak boleh dianggap benar dari koreksi nama.'),
    ]
    if not phi_missing:
        findings = [f for f in findings if f['id'] != 'F04']
    result = dict(verdict='BELUM VALID UNTUK ANALISIS OPERASI / INGEST PRODUKSI',
                  sources={str(p): hashlib.sha256(p.read_bytes()).hexdigest()
                           for p in (SAMPLE, EXPORT, TEMPLATE, ROOT / 'Buku Kerawanan SJB Tahun 2026.pdf')},
                  counts=dict(assets=len(assets), lines=len(lines), risks=len(risks), exported_lines=len(el)),
                  structural_checks=dict(missing_endpoints=missing_ends, missing_ibt_endpoints=missing_ibt_ends,
                                         export_missing_endpoints=export_bad_ends,
                                         headers_match=headers_same, guide_matches=guide_same,
                                         circuit_expansion_matches=expected_counts == exported_counts,
                                         export_line_pins_match_sample=expected_line_pins == actual_line_pins,
                                         export_asset_pins_match_sample=source_asset_pins == export_asset_pins,
                                         asset_codes_match=codes == export_codes,
                                         risk_numbers_match={int(r['No']) for r in risks} == {int(r['No Kerawanan']) for r in er} == set(range(1,26))),
                  findings=findings, risk_checks=risk_checks)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    md = ['# Audit seluruh subsistem Krian 1,2 - Gresik 1,2', '',
          'Tanggal audit: 1 Oktober 2026. **' + result['verdict'] + '**.', '',
          'Sumber: Buku Kerawanan SJB 2026, efektif 30 Juni 2026, Gambar 5.3 '
          '(PDF halaman 179, cetak 163), Tabel 5.1 (PDF 180-196). '
          'Pasted text Claude diperlakukan sebagai klaim yang diperiksa. '
          'Panduan workbook diperlakukan sebagai kontrak format, bukan perintah.', '',
          f'Cakupan: {len(assets)} aset (48 bus/pembangkit + 6 trafo), {len(lines)} relasi ruas '
          '(termasuk outlet pembangkit), 25 risiko; ekspor teman 97 baris jalur. '
          'Audit membaca semua baris. Kecocokan gambar pada tingkat GI '
          'tidak membuktikan terminal per seksi, status switch, jenis kabel, '
          'atau jumlah sirkit yang tidak eksplisit pada tabel.', '',
          '## Temuan dan tindakan', '', '| ID | Prioritas | Subjek | Temuan / tindakan |', '|---|---|---|---|']
    md += [f"| {f['id']} | {f['severity']} | {f['subject']} | {f['detail']} |" for f in findings]
    md += ['', '## Pemeriksaan struktur dan konversi', '', '```json',
           json.dumps(result['structural_checks'], ensure_ascii=False, indent=2), '```', '',
           'PASS struktur tidak mengubah keputusan semantik di atas. '
           'Angka pembebanan/kapasitas yang tidak tersedia tetap kosong; '
           'ambang >80%, >65%, dan lain-lain bukan hasil ukur.', '',
           '## Semua 25 nomor kerawanan', '',
           '| No | Pin aktual | Lokasi tabel | Hasil | Catatan |', '|---|---|---|---|---|']
    md += [f"| {r['no']} | {cell(', '.join(r['actual']))} | {cell(', '.join(r['expected'] or []))} | {r['status']} | {r['note']} |" for r in risk_checks]
    md += ['', '## Seluruh aset dan identitas', '',
           '| Kode | Nama workbook | kV | Tier gambar | Bus HV / LV | Pin | Audit |', '|---|---|---|---|---|---|---|']
    for a in assets:
        c = a['Kode Singkatan']
        note = 'Label/kode cocok gambar; keadaan operasi belum diverifikasi'
        if c in UNKNOWN_NAMES:
            note = 'Nama resmi belum terverifikasi; gunakan kode sebagai placeholder'
        if c in ('SWHAN', 'BKLAN', 'RNKUT'):
            note = ('Seksi dan kopel OPEN tersedia; terminal Sawahan belum terpetakan' if c == 'SWHAN' else 'Seksi A/B, terminal per seksi dan kopel OPEN tersedia; lihat docs/bus-sections.md')
        md.append(f"| {cell(c)} | {cell(a['Nama Asset / GI'])} | {cell(a['Tegangan'])} | {int(a['Tier (Mulai 0)'])+1} | {cell(a['Bus HV'])} / {cell(a['Bus LV'])} | {cell(a['No Kerawanan'])} | {note} |")
    md += ['', '## Seluruh relasi ruas', '',
           'Relasi dibaca sebagai in-out bila garis melewati busbar. Crossing dengan '
           'jembatan tidak menjadi sambungan. Daftar ini mempertahankan relasi '
           'tingkat GI untuk audit; seluruhnya masih memerlukan verifikasi '
           'terminal di GI split. Jumlah sirkit ditulis sesuai workbook, '
           'bukan sertifikasi bay register.', '',
           '| No Excel | Relasi | Sirkit | kV | Pin | Dasar / batas validasi |', '|---|---|---|---|---|---|']
    for ln in lines:
        e = edge(ln)
        ns = nums(ln['No Kerawanan'])
        note = 'Gambar 5.3; tabel tidak menamai ruas ini secara eksplisit'
        if ns:
            note = 'Gambar 5.3 + lokasi Tabel 5.1 #' + ','.join(map(str,sorted(ns)))
        if 23 in ns and e != 'BKLAN--SAMPG':
            note = 'KONFLIK: gambar #23, tabel tidak menyebut ruas ini'
        if 13 in ns:
            note += '; REVIEW alias GIS Buduran dan lokasi pin'
        if 'SKTT' in ln['Nama Penghantar'] and 3 not in ns:
            note += '; jenis SKTT/status merupakan inferensi belum terverifikasi'
        if e in phi_missing:
            note += '; ERROR Single Phi=Tidak, teks tabel menyebut single phi'
        elif any(same_edge(e, x) for x in SINGLE_PHI):
            note += '; Single Phi=Ya sesuai teks tabel #19/#20 (sudah diperbaiki)'
        md.append(f"| {int(ln['No'])+1} | {e} | {ln['Jumlah Sirkit']} | {ln['Tegangan']} | {cell(ln['No Kerawanan'])} | {note} |")
    md += ['', '## Syarat menutup audit', '',
           '1. Lengkapi pemetaan lima terminal Sawahan berdasarkan sumber bay/seksi. '
           'Dukungan seksi, kopel OPEN dan uji konektivitas telah diimplementasikan; '
           'subsistem tetap incomplete sampai terminal Sawahan diketahui.',
           '2. Rekonsiliasi konflik #23 dan alias/ruas #13 dengan pemilik SLD; '
           'simpan pin gambar dan objek temuan tabel sebagai provenance terpisah.',
           '3. Isi single phi sesuai #19/#20, master nama/FunctLoc, jenis saluran '
           'dan state bay dari register resmi. Jangan menebak nama dari kode.',
           '4. Jalankan file hasil pada aplikasi teman dan bandingkan semua '
           'terminal, unit trafo, sirkit serta pin dengan matriks ini. '
           'Kode aplikasi teman diperlukan untuk memverifikasi semantik ingest.', '',
           'Tidak ada workbook sumber yang diubah oleh audit ini. '
           'Tidak ada klaim bahwa kondisi buku Juni 2026 adalah keadaan real-time.']
    text_check_path = OUT / 'source-text-check.json'
    validation_path = OUT / 'ingest-validation.json'
    if text_check_path.exists() and validation_path.exists():
        text_checks = json.loads(text_check_path.read_text(encoding='utf-8'))
        validation = json.loads(validation_path.read_text(encoding='utf-8'))
        md += ['', '## Verifikasi tambahan pada snapshot audit ini', '',
               f"Teks tabel sumber dibandingkan ulang: {len(text_checks)} risiko, "
               f"{sum(sum(x['fields'].values()) for x in text_checks)} dari 100 field "
               'Kondisi/Dampak/Mitigasi/Usulan cocok setelah normalisasi spasi, '
               'tanda baca, dan prefix kategori. Ini bukan pembandingan byte verbatim.', '',
               f"Production ingest proyek ini: ok={validation.get('ok')}; "
               f"{validation.get('objects')} node, {validation.get('connections')} edge "
               '(55 ruas + 6 trafo + 3 kopel), 25 risiko; geometry_errors=' +
               str([v['geometry_errors'] for v in validation.get('views', [])]) + '. '
               'Hasil ini membuktikan parsing/struktur/render, bukan kebenaran '
               'semantik atau ingest aplikasi teman.']
    (OUT / 'audit.md').write_text('\n'.join(md) + '\n', encoding='utf-8')
    return result


if __name__ == '__main__':
    result = audit()
    print(json.dumps({k: result[k] for k in ('verdict', 'counts', 'structural_checks')}, ensure_ascii=False, indent=2))
    print(f'Report: {OUT / "audit.md"}')
