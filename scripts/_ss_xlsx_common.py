"""Shared writer for the per-subsystem ingest workbooks under scripts/make_ss_*.py.

The workbook is the PLN subsystem template the /ingest parser reads
(app/services/ingest_parser.py): sheets Info, Views (optional multi-SLD),
Gardu_Induk_dan_Aset, Jalur_Transmisi, Bay (optional), Data_Kerawanan_Detail.

One SS = one spec dict. Relations come from the book SLD (Gambar 2.x), traced
as-is; a ruas with confidence < 0.9 lands NEEDS_REVIEW in the engine. GITET /
GI codes follow the book's labels -- they are cosmetic; the relations are what
must be right.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

SAMPLES = Path(__file__).resolve().parents[1] / "samples"

ASSET_HEADER = ["No", "Nama Asset / GI", "Kode Singkatan", "Tipe Asset",
                "Tier (Mulai 0)", "Tegangan", "No IBT", "Bus 150 kV",
                "Jumlah Trafo", "Jumlah Kapasitor", "Catatan Simbol",
                "Jumlah Sirkit Bay", "Status Operasi", "Role", "Status Kerawanan",
                "No Kerawanan", "Wilayah", "Sudut Pandang", "Latitude", "Longitude"]
LINE_HEADER = ["No", "No Kerawanan", "Nama Penghantar", "Dari GI", "Ke GI",
               "Tegangan", "Panjang Saluran (km)", "Jumlah Sirkit",
               "Status Operasi", "Tingkat Kerawanan", "Pembebanan Sirkit 1 (%)",
               "Pembebanan Sirkit 2 (%)", "Koridor / Wilayah", "Tier Dari",
               "Tier Ke", "Single Phi", "Sudut Pandang"]
BAY_HEADER = ["No", "Kode GI", "Nama GI", "Feeder (GI Induk)", "Tegangan",
              "Jumlah Sirkit", "Status Operasi", "No Kerawanan", "Sudut Pandang"]
RISK_HEADER = ["No", "UIT", "Kategori Kontingensi", "Kondisi / Permasalahan",
               "Dampak", "Mitigasi", "Usulan / Solusi"]


def _bold(ws):
    for c in ws[1]:
        c.font = Font(bold=True)


def _rawan(no_kerawanan, status):
    if no_kerawanan:
        return "Sangat Rawan"
    if status and status.lower() in ("rencana", "planned"):
        return "Rawan"
    return "Normal"


def build_workbook(spec: dict) -> Path:
    """spec keys:
        code, name, apb, wilayah, source_ref
        views:   [ (view_key, view_name, "GI;GI;..", page) ]   (optional; omit -> single SLD)
        assets:  [ dict(code, name, type, tier, kv=150, ibt=None, bus150=None,
                        status="Beroperasi", kerawanan=None, views=None) ]
            type: "Busbar GITET" | "Busbar GI" | "Busbar GIS" | "Pembangkit" |
                  "IBT 3-Winding" | "Bay"
            status: "Beroperasi" | "Rencana" | "Belum Operasi"
        lines:   [ dict(fr, to, name, kv=150, sirkit=2, status="Beroperasi",
                        kerawanan=None, koridor=None, tier_fr=None, tier_to=None,
                        views=None) ]
        bays:    [ (gi_code, gi_name, feeder_code, kerawanan_or_None
                    [, circuit_count[, status]]) ]
        risks:   [ dict(no, uit="JBB", category="N-1", kondisi, dampak,
                        mitigasi, usulan) ]
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wil = spec.get("wilayah", "DKI Jakarta")
    views = spec.get("views") or []
    default_view = views[0][0] if views else None

    ws = wb.create_sheet("Info")
    ws.append(["Kunci", "Nilai"]); _bold(ws)
    for k, v in [("Kode Subsistem", spec["code"]),
                 ("Nama Subsistem", spec["name"]),
                 ("APB", spec.get("apb", "UP2B Jakarta & Banten"))]:
        ws.append([k, v])
    if spec.get("source_ref"):
        ws.append(["Sumber", spec["source_ref"]])

    if views:
        ws = wb.create_sheet("Views")
        ws.append(["Kode View", "Nama View", "Sumber Tier-1 (kode GI, pisah ;)",
                   "Halaman Buku"]); _bold(ws)
        for (vk, vn, src, page) in views:
            ws.append([vk, vn, src, page])

    ws = wb.create_sheet("Gardu_Induk_dan_Aset")
    ws.append(ASSET_HEADER); _bold(ws)
    for i, a in enumerate(spec["assets"], 1):
        kno = a.get("kerawanan")
        vw = ";".join(a["views"]) if a.get("views") else default_view
        ws.append([i, a["name"], a["code"], a["type"],
                   a["tier"] - 1 if a["tier"] else a["tier"],
                   a.get("kv", "150 kV") if isinstance(a.get("kv"), str)
                   else f"{a.get('kv', 150)} kV",
                   a.get("ibt", ""), a.get("bus150", ""),
                   a.get("trafo", ""), a.get("kapasitor", ""), a.get("simbol", ""),
                   a.get("bay_sirkit", ""), a.get("status", "Beroperasi"), a.get("role", ""),
                   "Rawan" if kno else "Normal", kno or "", wil, vw or ""])
        ws.cell(ws.max_row, len(ASSET_HEADER) - 1, a.get("latitude"))
        ws.cell(ws.max_row, len(ASSET_HEADER), a.get("longitude"))

    ws = wb.create_sheet("Jalur_Transmisi")
    ws.append(LINE_HEADER); _bold(ws)
    for i, ln in enumerate(spec["lines"], 1):
        kno = ln.get("kerawanan")
        status = ln.get("status", "Beroperasi")
        vw = ";".join(ln["views"]) if ln.get("views") else default_view
        ws.append([i, kno or "", ln["name"], ln["fr"], ln["to"],
                   f"{ln.get('kv', 150)} kV" if not isinstance(ln.get("kv"), str)
                   else ln["kv"],
                   "", ln.get("sirkit", 2), status, _rawan(kno, status),
                   "", "", ln.get("koridor", wil),
                   ln.get("tier_fr", ""), ln.get("tier_to", ""),
                   "Ya" if ln.get("single_phi") or "single phi" in ln["name"].lower() else "Tidak",
                   vw or ""])

    if spec.get("bays"):
        ws = wb.create_sheet("Bay")
        ws.append(BAY_HEADER); _bold(ws)
        for i, bay in enumerate(spec["bays"], 1):
            gc, gn, fd, kno, *extra = bay
            circuit_count = extra[0] if len(extra) > 0 else 1
            status = extra[1] if len(extra) > 1 else "Beroperasi"
            views_for_bay = extra[2] if len(extra) > 2 else ""
            if isinstance(views_for_bay, (list, tuple)):
                views_for_bay = ";".join(views_for_bay)
            ws.append([i, gc, gn, fd, "150 kV", circuit_count, status, kno or "", views_for_bay])

    ws = wb.create_sheet("Data_Kerawanan_Detail")
    ws.append(RISK_HEADER); _bold(ws)
    for r in spec["risks"]:
        cat = r.get("category", "N-1")
        ws.append([r["no"], r.get("uit", "JBB"), cat,
                   f"[{cat}] {r['kondisi']}", r.get("dampak", ""),
                   r.get("mitigasi", ""), r.get("usulan", "")])

    SAMPLES.mkdir(exist_ok=True)
    out = SAMPLES / f"{spec['code'].lower()}_ingest.xlsx"
    wb.save(out)
    n_kit = sum(1 for a in spec["assets"] if a["type"] == "Pembangkit")
    print(f"wrote {out.name}: {len(spec['assets'])} aset ({n_kit} KIT), "
          f"{len(spec['lines'])} ruas, {len(spec['risks'])} kerawanan, "
          f"{len(views)} view")
    return out
