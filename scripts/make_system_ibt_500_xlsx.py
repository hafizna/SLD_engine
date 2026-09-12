"""Standalone 500 kV IBT risk-map workbook, PDF figure 1.5 / table 1.2.

Does not alter SS fixtures. Review inventory is explicitly separate from map
topology: incomplete AI-extracted SS endpoints must not change the backbone.
"""
from pathlib import Path
from collections import defaultdict
import sys
import hashlib
import re
import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.services.ingest_parser import parse_upload
from _kerawanan_tables import clean, _strip_chrome

OUT = ROOT / 'samples/system_ibt_500_ingest.xlsx'
TARGETS = {
    1: [('SRLYA','1,2')], 2: [('SRLYA','1,2'),('CLGON','4')], 3:[('CLGON','lihat kondisi')],
    4:[('KMBNG','1,2')],5:[('CWANG','2,3'),('DEPOK','1')],6:[('BKASI','1,3'),('CIBNG','3')],
    7:[('CBATU','3,4')],8:[('MDCAN','1,2')],9:[('CRATA','1,2,3')],10:[('CBATU','1,2')],
    11:[('DLTMS','3,4')],12:[('DLTMS','1,2')],13:[('BDSLN','1,2')],14:[('UBRNG','1,2')],
    15:[('TSMYA','1,2')],16:[('TMBUN','1,2')],17:[('TJATI','1,2'),('UNGRN','3')],
    18:[('TJATI','1,2'),('UNGRN','3')],19:[('UNGRN','3'),('TJATI','1,2')],20:[('UNGRN','3')],
    21:[('UNGRN','1,2')],22:[('UNGRN','1,2,3')],23:[('PEDAN','1,2')],24:[('PEDAN','3,4')],
    25:[('PEDAN','1,2,3,4')],26:[('KSGHN','1,2')],27:[('PMLNG','1,2')],28:[('BYOLI','1,2')],
    29:[('UNGRN','3'),('BYOLI','1,2')],30:[('KRIAN','1,2')],31:[('GRSIK','1,2')],
    32:[('KRIAN','3,4,5,6')],33:[('NBANG','lihat kondisi')],34:[('KDIRI','1,2')],
    35:[('KDIRI','3,4')],36:[('GRATI','1,2,3')],37:[('PITON','2')],38:[('PITON','1,3')],
}

def extract_risks():
    records, current = [], None
    with pdfplumber.open(ROOT / 'Buku Kerawanan SJB Tahun 2026.pdf') as pdf:
        for page in range(41,65):
            for table in pdf.pages[page-1].extract_tables():
                for row in table:
                    if len(row) != 7: continue
                    cells = [clean(v) for v in row]
                    if cells[0].isdigit() and cells[1] in ('JBB','JBT','JBM'):
                        current = {'seq':int(cells[0]),'page':page,'cells':cells}
                        records.append(current)
                    elif current and not cells[0] and cells[1] not in ('UIT','UIT / UNIT'):
                        for i in range(2,7):
                            if cells[i]: current['cells'][i] += ' ' + cells[i]
    assert [r['seq'] for r in records] == list(range(1,39)), 'Incomplete table 1.2'
    for r in records: r['cells'][2:] = [_strip_chrome(x) for x in r['cells'][2:]]
    return records

def main():
    risks = extract_risks()
    book = openpyxl.load_workbook(ROOT/'samples/backbone_500_ingest.xlsx')
    info = book['Info']
    info.delete_rows(1,info.max_row)
    for row in [
        ('Kunci','Nilai'),('Kode Subsistem','SYSTEM_IBT_500'),
        ('Nama Subsistem','Sistem Jamali 500 kV - Kerawanan IBT 500/150 kV'),
        ('APB','UIP2B JAMALI'),('Rule Profile','IBT_500_150'),
        ('Sumber risiko','Buku Kerawanan SJB 2026, Gambar 1.5 (PDF40), Tabel 1.2 (PDF41-64)'),
        ('Basis jaringan','backbone_500_ingest.xlsx; topologi 500 kV untuk peta, bukan jaringan SS'),
        ('Status','DRAFT UNTUK REVIEW; belum dipublikasikan'),
        ('Pin peta','Pin pada GITET utama; unit dan seluruh target ada di Target_Kerawanan_IBT'),
        ('Batasan target','Target tambahan tersimpan untuk review; parser saat ini memakai satu pin utama per risiko'),
        ('Inventaris','IBT_Unit_Review berasal dari workbook SS; bukan data final. Tidak menambah bus LV pada peta 500 kV'),
        ('Kategori','Kategori pada tabel detail adalah kandidat dari teks kondisi; review sebelum publikasi'),
        ('Cara review','Review hubungan 500 kV di Jalur_Transmisi, target risiko, lalu unit/endpoint pada IBT_Unit_Review'),
    ]: info.append(row)
    views = book['Views']
    views.cell(2,1,'IBT500');views.cell(2,2,'Peta Kerawanan IBT 500/150 kV');views.cell(2,4,24)
    pins=defaultdict(list)
    for seq,targets in TARGETS.items(): pins[targets[0][0]].append(seq)
    assets=book['Gardu_Induk_dan_Aset']
    header={c.value:c.column for c in assets[1]}
    codes=set()
    for row in range(2,assets.max_row+1):
        code=assets.cell(row,header['Kode Singkatan']).value; codes.add(code)
        seqs=pins.get(code,[])
        assets.cell(row,header['No Kerawanan']).value=';'.join(map(str,seqs)) or None
        assets.cell(row,header['Status Kerawanan']).value='Rawan' if seqs else 'Normal'
        assets.cell(row,header['Sudut Pandang']).value='IBT500'
    assert all(code in codes for targets in TARGETS.values() for code,_ in targets)
    lines=book['Jalur_Transmisi']; lh={c.value:c.column for c in lines[1]}
    for row in range(2,lines.max_row+1):
        lines.cell(row,lh['No Kerawanan']).value=None
        lines.cell(row,lh['Tingkat Kerawanan']).value='Normal'
        lines.cell(row,lh['Sudut Pandang']).value='IBT500'
    risk=book['Data_Kerawanan_Detail'];risk.delete_rows(1,risk.max_row)
    risk.append(['No','UIT','Kategori Kontingensi','Kondisi / Permasalahan','Dampak','Mitigasi','Usulan / Solusi','Subsistem sumber','Halaman PDF','Status Review'])
    target=book.create_sheet('Target_Kerawanan_IBT')
    target.append(['No Risiko','GITET','Unit disebut sumber','Peran pin','Halaman PDF','Status Review'])
    for r in risks:
        no,uit,ss,condition,impact,mitigation,followup=r['cells']
        categories=list(dict.fromkeys(re.findall(r'N\s*-\s*[12](?:\s*-\s*[12])?',condition)))
        category=';'.join(re.sub(r'\s+','',s) for s in categories) or 'BELUM_DITETAPKAN'
        risk.append([int(no),uit,category,condition,impact,mitigation,followup,ss,r['page'],'PERLU REVIEW'])
        for idx,(code,units) in enumerate(TARGETS[int(no)]):
            target.append([int(no),code,units,'PIN UTAMA' if idx==0 else 'OBJEK TERKAIT',r['page'],'PERLU REVIEW'])
    inventory=book.create_sheet('IBT_Unit_Review')
    inventory.append(['Workbook SS','SHA256 sumber','Kode SS','Bus HV','kV HV','Unit','Bus LV','kV LV','View SS','Status operasi sumber','Status Review'])
    for path in sorted((ROOT/'samples').glob('ss_*_ingest.xlsx')):
        blob=path.read_bytes(); payload=parse_upload(blob,path.name)
        nodes={n['external_key']:n for n in payload['objects']}
        for e in payload['connections']:
            hv=nodes[e['from_external_key']];lv=nodes[e['to_external_key']]
            if e['relation_type']!='IBT_LINK' or hv['voltage_hv_kv']!=500: continue
            inventory.append([path.name,hashlib.sha256(blob).hexdigest(),payload['subsystem']['code'],
                e['from_external_key'],500,e['unit_no'],e['to_external_key'],lv['voltage_hv_kv'],
                ';'.join(e['view_keys']),e['status_hint'],'PERLU REVIEW; cek unit yang digabung pada sumber'])
    for ws in book:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:
            c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='0047AB')
            c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[1].height=32
        for col in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width=24
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    info.column_dimensions['A'].width=26;info.column_dimensions['B'].width=110
    for row in range(2,info.max_row+1): info.row_dimensions[row].height=32
    for col in 'DEFG': risk.column_dimensions[col].width=68
    for row in range(2,risk.max_row+1):
        longest=max(len(str(risk.cell(row,col).value or '')) for col in range(4,8))
        risk.row_dimensions[row].height=min(409,max(80,(math.ceil(longest/60)+2)*15))
    inventory.column_dimensions['A'].width=43;inventory.column_dimensions['B'].width=68
    inventory.column_dimensions['K'].width=58
    for ws in (assets,lines,inventory,target):
        for row in range(2,ws.max_row+1): ws.row_dimensions[row].height=48
    book.save(OUT)
    print(f'{OUT}: {len(risks)} risks, {target.max_row-1} target rows, {inventory.max_row-1} SS IBT rows')

if __name__=='__main__': main()
