"""Rebuild ingest samples from the user's reviewed Excel sources.

Keep audit prose as evidence, not as machine-readable view membership.
The original workbooks under samples/sources are never modified.
"""
from pathlib import Path
import json

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "SS_LBK": "JBB_SS_LBK_single_view_v1.xlsx",
    "SS_GUCL": "JBB_SS_GUCL_manual_audit_v2 (1).xlsx",
    "SS_PRBC": "JBB_SS_PBRC_single_view_v1.xlsx",
}
LEGACY_NAMES = json.loads(
    (ROOT / "samples" / "sources" / "jakban_legacy_names.json").read_text(encoding="utf-8")
)


def _headers(ws):
    return {str(cell.value).strip(): index + 1
            for index, cell in enumerate(ws[1]) if cell.value is not None}


def _column(headers, *names):
    for name in names:
        for header, index in headers.items():
            if header.casefold() == name.casefold():
                return index
    return None


def _append_record(ws, values):
    headers = _headers(ws)
    row = [None] * max(headers.values())
    for key, value in values.items():
        index = _column(headers, key)
        if index:
            row[index - 1] = value
    ws.append(row)


def _row_value(ws, row, *names):
    headers = _headers(ws)
    index = _column(headers, *names)
    return ws.cell(row, index).value if index else None


def _next_no(ws):
    values = []
    for row in range(2, ws.max_row + 1):
        value = _row_value(ws, row, "No")
        try:
            values.append(int(value))
        except (TypeError, ValueError):
            pass
    return max(values, default=0) + 1


def _legacy_name(subsystem, code, asset_type, voltage, fallback):
    """Return the old workbook's clearer label when the code is stable.

    The code is the join key. Voltage disambiguates the old workbook's shared
    GITET/GI codes; the ``GITET_`` fallback handles the explicit split used by
    the reviewed PBRC model. A type mismatch is tolerated because the reviewed
    book intentionally corrected a few GIS/GI classifications.
    """
    names = LEGACY_NAMES.get(subsystem, {})
    wanted_type = str(asset_type or "").strip()
    wanted_voltage = str(voltage or "").strip()
    exact = f"{code}|{wanted_type}|{wanted_voltage}"
    if exact in names:
        return names[exact]
    base = code[6:] if code.startswith("GITET_") else code
    for key, value in names.items():
        raw_code, raw_type, raw_voltage = key.split("|", 2)
        if raw_code == base and raw_voltage == wanted_voltage:
            return value
    for key, value in names.items():
        raw_code, _raw_type, raw_voltage = key.split("|", 2)
        if raw_code == base and (not wanted_voltage or raw_voltage == wanted_voltage):
            return value
    # A boundary Bay can intentionally represent the other voltage side of a
    # shared old code (for example NCKUPA). Keep its established site label
    # even when the old asset row used only the 500 kV spelling.
    for key, value in names.items():
        raw_code, _raw_type, _raw_voltage = key.split("|", 2)
        if raw_code == base:
            return value
    return fallback


def _restore_legacy_names(wb, subsystem):
    """Enrich code-only reviewed rows with stable labels from the old book."""
    ws = wb["Gardu_Induk_dan_Aset"]
    headers = _headers(ws)
    code_col = _column(headers, "Kode Singkatan", "Kode", "Code")
    name_col = _column(headers, "Nama Asset / GI", "Nama Asset", "Nama GI", "Name")
    type_col = _column(headers, "Tipe Asset", "Tipe", "Type")
    voltage_col = _column(headers, "Tegangan", "Voltage")
    if not all((code_col, name_col)):
        return
    for row in range(2, ws.max_row + 1):
        raw_code = ws.cell(row, code_col).value
        if not raw_code:
            continue
        current = ws.cell(row, name_col).value
        label = _legacy_name(
            subsystem,
            str(raw_code).strip(),
            ws.cell(row, type_col).value if type_col else None,
            ws.cell(row, voltage_col).value if voltage_col else None,
            current,
        )
        if label and label != current:
            ws.cell(row, name_col).value = label


def _augment_boundary_evidence(wb, code):
    """Promote explicit source/stub evidence into the ingest model.

    The reviewed books keep their audit sheets as evidence. A small, explicit
    subset is also copied into the machine-readable sheets: source generators,
    500/150 GITET-to-bus links, and boundary stubs whose feeder is unambiguous.
    Ambiguous gray corridors remain evidence-only.
    """
    asset = wb["Gardu_Induk_dan_Aset"]
    line = wb["Jalur_Transmisi"]
    bay = wb["Bay"]
    asset_headers = _headers(asset)
    asset_codes = {
        str(_row_value(asset, row, "Kode Singkatan", "Kode", "Code")).strip()
        for row in range(2, asset.max_row + 1)
        if _row_value(asset, row, "Kode Singkatan", "Kode", "Code")
    }
    line_pairs = {
        (str(_row_value(line, row, "Dari GI", "From")).strip(),
         str(_row_value(line, row, "Ke GI", "To")).strip())
        for row in range(2, line.max_row + 1)
        if _row_value(line, row, "Dari GI", "From")
        and _row_value(line, row, "Ke GI", "To")
    }
    bay_pairs = {
        (str(_row_value(bay, row, "Kode GI", "Kode", "Code")).strip(),
         str(_row_value(bay, row, "Feeder (GI Induk)", "Feeder", "GI Induk")).strip())
        for row in range(2, bay.max_row + 1)
        if _row_value(bay, row, "Kode GI", "Kode", "Code")
        and _row_value(bay, row, "Feeder (GI Induk)", "Feeder", "GI Induk")
    }

    def add_asset(asset_code, name, asset_type, tier, voltage, status="Beroperasi",
                  role=None, note=None, bus_hv=None, bus_lv=None, unit=None):
        if asset_code in asset_codes:
            return
        name = _legacy_name(code, asset_code, asset_type, voltage, name)
        _append_record(asset, {
            "No": _next_no(asset), "Nama Asset / GI": name,
            "Kode Singkatan": asset_code, "Tipe Asset": asset_type,
            "Tier (Mulai 0)": tier, "Tegangan": voltage,
            "No IBT": unit, "Bus HV": bus_hv, "Bus LV": bus_lv,
            "Jumlah Trafo": 1 if "IBT" in asset_type.upper() else None,
            "Status Operasi": status, "Role": role, "Catatan Simbol": note,
            "Wilayah": "Jakarta & Banten" if code == "SS_PRBC" else "Banten",
        })
        asset_codes.add(asset_code)

    def add_line(name, fr, to, voltage="150 kV", circuits=1, tier_fr=1, tier_to=1,
                 status="Beroperasi"):
        if (fr, to) in line_pairs or (to, fr) in line_pairs:
            return
        _append_record(line, {
            "No": _next_no(line), "Nama Penghantar": name,
            "Dari GI": fr, "Ke GI": to, "Tegangan": voltage,
            "Jumlah Sirkit": circuits, "Status Operasi": status,
            "Tingkat Kerawanan": "Normal", "Koridor / Wilayah": "Jakarta & Banten",
            "Tier Dari": tier_fr, "Tier Ke": tier_to, "Single Phi": "Tidak",
        })
        line_pairs.add((fr, to))

    def add_bay(stub, name, feeder, kind="Boundary / external", status="Belum Operasi"):
        if (stub, feeder) in bay_pairs:
            return
        name = _legacy_name(code, stub, "Busbar GI", "150 kV", name)
        _append_record(bay, {
            "No": _next_no(bay), "Kode GI": stub, "Nama GI": name,
            "Feeder (GI Induk)": feeder, "Jenis": kind, "Tegangan": "150 kV",
            "Jumlah Sirkit": 1, "Status Operasi": status,
            "Catatan Sudut Pandang Sumber": "Boundary_External; source evidence",
        })
        bay_pairs.add((stub, feeder))

    if code == "SS_PRBC":
        # The source uses "Busbar GITET" for 150 kV endpoint buses. The
        # Boundary_External sheet identifies the actual 500 kV GITETs.
        for bus in ("BKASI", "MTWAR", "CWBRU"):
            for row in range(2, asset.max_row + 1):
                if _row_value(asset, row, "Kode Singkatan", "Kode", "Code") == bus:
                    typ = _column(asset_headers, "Tipe Asset", "Tipe", "Type")
                    kv = _row_value(asset, row, "Tegangan", "Voltage")
                    if typ and str(kv).startswith("150"):
                        asset.cell(row, typ).value = "Busbar GI"
                    break
        for gitet, bus, units in (
            ("GITET_BKASI", "BKASI", ("2", "4")),
            ("GITET_MTWAR", "MTWAR", ("1", "2")),
            ("GITET_CWBRU", "CWBRU", ("1",)),
        ):
            add_asset(gitet, f"{gitet.replace('GITET_', 'GITET ')}", "Busbar GITET", 0,
                      "500 kV", role="SOURCE", note="Boundary_External: 500/150 kV source")
            for unit in units:
                add_asset(f"IBT {unit} {bus}", f"IBT {unit} {bus}", "IBT 3-Winding", 1,
                          "500/150 kV", role="SOURCE_BOUNDARY", bus_hv=gitet,
                          bus_lv=bus, unit=unit)
        for kit, name, bus in (
            ("KIT_MKR_ST30", "PLTGU Muarakarang ST 3.0", "MKLMA"),
            ("KIT_PRIOK_B12", "PLTGU Priok Blok 1 & 2", "PRBRT"),
            ("KIT_PRIOK_B3", "PLTGU Priok Blok 3", "PRTRU"),
        ):
            add_asset(kit, name, "Pembangkit", 0, "150 kV", role="SOURCE",
                      note="Boundary_External: Source / generator")
            add_line(f"Outlet {name}", kit, bus, circuits=1)
        for stub in ("PDKLP", "SKTNI", "SMRCN"):
            add_bay(stub, f"{stub} (boundary stub di BKASI)", "BKASI")

    elif code == "SS_LBK":
        # These are explicit black/boundary stubs in the review sheet. Keep
        # uncertain names visible as external context without promoting them
        # into the red core network.
        for stub, name, feeder, status in (
            ("NCKUPA", "GITET New Cikupa (boundary)", "CKUPA", "Belum Operasi"),
            ("CKNN?", "CKNN? (label raster)", "DLRA", "Belum Operasi"),
            ("TGBRU3", "Tangerang Baru 3 (boundary)", "SUJYA", "Belum Operasi"),
            ("JTKBR", "Jatake Baru (boundary)", "JTAKE", "Belum Operasi"),
            ("BSH", "BSH (KTT / customer asset)", "CKDRU", "Milik Pelanggan"),
        ):
            add_bay(stub, name, feeder, "KTT / customer" if stub == "BSH" else "Boundary / external", status)

    elif code == "SS_GUCL":
        # The gray external chain is explicit in the audit sheet. Promote only
        # the unambiguous parent-child chain; KOPO remains unresolved evidence.
        for stub, name, feeder in (
            ("SARAN4", "SARAN4 (boundary external)", "SRANG"),
            ("RGKOT", "RGKOT (boundary external)", "SARAN4"),
            ("BUNAR", "BUNAR (boundary external)", "RGKOT"),
            ("KRACAK", "KRACAK (boundary spur)", "BUNAR"),
        ):
            add_bay(stub, name, feeder)


# The reviewed LBK and PBRC books carry no risk table, so the dashboard showed
# both subsystems with 0 kerawanan. Their Tabel 2.3 / 2.8 rows still live in
# the historical SPEC of each builder; they are re-attached to the reviewed
# topology here. A pin marks where the finding sits -- the objects its Kondisi
# names (the Sumatera rule, user 2026-09-23) -- never the GIs it knocks out.
# Where the reviewed topology no longer has the named object, the nearest one
# that is drawn stands in, and the note says so.
RISK_PINS = {
    "SS_LBK": {
        # IBT-1,2 Kembangan: the reviewed book models Kembangan's 150 kV bus
        # only, with no GITET or IBT rows -- the pin stands on that bus.
        1: [("asset", "KMBGN")],
        2: [("line", "KMBGN", "NSYAN")],
        3: [("line", "PSKBR", "GJTGL"), ("line", "PSKMS", "GJTGL")],
        4: [("line", "CKUPA", "JTAKE"), ("line", "ILKNG", "TGBRU")],
        # Durikosambi - Cengkareng: Durikosambi is SS Muarakarang's, and the
        # reviewed book does not draw the ruas; Cengkareng is its end here.
        5: [("asset", "CNKNG")],
        6: [("asset", "SNYAN"), ("line", "NSYAN", "SNYAN")],
    },
    "SS_PRBC": {
        # "interconnector-1&2 Priok Timur Lama arah Priok Barat": the reviewed
        # book has no direct Timur Lama - Barat ruas; the interconnector drawn
        # at Priok Timur is Timur Baru - Timur Lama.
        1: [("line", "PRTMR", "PRTRU")],
        2: [("line", "PRBRT", "PLPNG40"), ("line", "PRTMR", "PLPNG20")],
        3: [("line", "GDPLA", "MGRAI"), ("line", "MGRAI", "DKTAS")],
        4: [("line", "PGLNG", "RATER")],
        5: [("asset", "KIT_PRIOK_B12"), ("asset", "KIT_PRIOK_B3")],
        6: [("asset", "PRBRT"), ("asset", "PRTMR"), ("asset", "PRTRU"),
            ("asset", "PLPNG20"), ("asset", "PLPNG40"), ("asset", "PGSAN")],
        7: [("asset", "GDPLA")],
    },
}


def _legacy_risks(code):
    """Tabel rows from the builder's historical SPEC (Buku Kerawanan SJB 2026)."""
    import importlib
    import sys
    scripts = str(ROOT / "scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    return importlib.import_module(f"make_{code.lower()}_xlsx").RISKS


def _mark(ws, row, number, rawan):
    headers = _headers(ws)
    col = _column(headers, "No Kerawanan", "No. Kerawanan")
    if col is None:
        col = max(headers.values()) + 1
        ws.cell(1, col).value = "No Kerawanan"
    current = ws.cell(row, col).value
    numbers = [n for n in str(current or "").replace(",", ";").split(";") if n.strip()]
    if str(number) not in numbers:
        numbers.append(str(number))
    ws.cell(row, col).value = ";".join(numbers)
    status = _column(headers, "Status Kerawanan", "Tingkat Kerawanan")
    if status:
        ws.cell(row, status).value = rawan


def _restore_risks(wb, code):
    pins = RISK_PINS.get(code)
    risk_ws = wb["Data_Kerawanan_Detail"]
    has_rows = any(_row_value(risk_ws, row, "No") not in (None, "")
                   for row in range(2, risk_ws.max_row + 1))
    if not pins or has_rows:
        return
    # drop the blank placeholder rows the reviewed books ship with
    for row in range(risk_ws.max_row, 1, -1):
        if all(cell.value in (None, "") for cell in risk_ws[row]):
            risk_ws.delete_rows(row)
    for r in _legacy_risks(code):
        cat = r.get("category", "N-1")
        _append_record(risk_ws, {
            "No": r["no"], "UIT": r.get("uit", "JBB"), "Kategori Kontingensi": cat,
            "Kondisi / Permasalahan": f"[{cat}] {r['kondisi']}",
            "Dampak": r.get("dampak", ""), "Mitigasi": r.get("mitigasi", ""),
            "Usulan / Solusi": r.get("usulan", ""),
        })
    asset, line = wb["Gardu_Induk_dan_Aset"], wb["Jalur_Transmisi"]
    for number, targets in pins.items():
        for target in targets:
            if target[0] == "asset":
                rows = [row for row in range(2, asset.max_row + 1)
                        if _row_value(asset, row, "Kode Singkatan", "Kode", "Code") == target[1]]
                ws, rawan = asset, "Rawan"
            else:
                pair = {target[1], target[2]}
                rows = [row for row in range(2, line.max_row + 1)
                        if {_row_value(line, row, "Dari GI", "From"),
                            _row_value(line, row, "Ke GI", "To")} == pair]
                ws, rawan = line, "Sangat Rawan"
            if not rows:
                raise ValueError(f"{code} kerawanan #{number}: {target[1:]} tidak ada di workbook revisi")
            for row in rows:
                _mark(ws, row, number, rawan)
    info = wb["Info"]
    if not any(row[0].value == "Multi Pin" for row in info.iter_rows()):
        info.append(["Multi Pin", "Ya"])


def build_reviewed(code: str, output_dir: Path | None = None) -> Path:
    source = ROOT / "samples" / "sources" / SOURCES[code]
    wb = openpyxl.load_workbook(source)
    for row in wb["Info"]:
        if row[0].value == "Kode Subsistem":
            row[1].value = code
    for name in ("Gardu_Induk_dan_Aset", "Jalur_Transmisi", "Bay"):
        ws = wb[name]
        headers = [c.value for c in ws[1]]
        if "Sudut Pandang" not in headers:
            continue
        column = headers.index("Sudut Pandang") + 1
        # Preserve the source annotation in place. The parser ignores this
        # explicitly named evidence column; all rows belong to the FULL view.
        ws.cell(1, column).value = "Catatan Sudut Pandang Sumber"
    if "Views" in wb:
        del wb["Views"]
    ws = wb.create_sheet("Views", 1)
    ws.append(["Kode View", "Nama View", "Sumber Tier-1 (kode GI, pisah ;)", "Halaman Buku"])
    roots = {"SS_LBK": "KMBGN;NBRJA;ILKNG",
             "SS_GUCL": "CLBRU;LBUAN",
             "SS_PRBC": "MKLMA;PRBRT;PRTMR;PRTRU;BKASI;MTWAR;CWBRU"}
    ws.append(["FULL", "SLD lengkap", roots[code], None])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in {"A": 16, "B": 24, "C": 65, "D": 18}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    _restore_legacy_names(wb, code)
    _augment_boundary_evidence(wb, code)
    _restore_risks(wb, code)
    output_dir = Path(output_dir) if output_dir else ROOT / "samples"
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / f"{code.lower()}_ingest.xlsx"
    wb.save(target)
    return target
