"""Apply reviewed MKBRU/DMGOT corrections (user 2026-09-12; PDF pp91-92).

Preserves unrelated rows and workbook formatting. Safe to rerun.
"""
from pathlib import Path
from copy import copy
import openpyxl

PATH = Path(__file__).resolve().parents[1] / 'samples/ss_muarakarang_durikosambi_ingest.xlsx'


def main():
    book = openpyxl.load_workbook(PATH)
    assets = book['Gardu_Induk_dan_Aset']
    for header in ('Bus HV', 'Bus LV'):
        if header not in [c.value for c in assets[1]]:
            assets.cell(1, assets.max_column + 1, header)
    def rows(sheet):
        headers = [c.value for c in sheet[1]]
        return [(i, dict(zip(headers, values))) for i, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), 2)]
    def setrow(sheet, index, values):
        headers = [c.value for c in sheet[1]]
        for key, value in values.items():
            sheet.cell(index, headers.index(key) + 1).value = value
    def append(sheet, values):
        index = sheet.max_row + 1
        setrow(sheet, index, values)
        for c in sheet[index]:
            c._style = copy(sheet.cell(index - 1, c.column)._style)

    for i, r in reversed(rows(assets)):
        if r['Kode Singkatan'] == 'GIS MKBRU' and r['Tipe Asset'] == 'Busbar GITET':
            assets.delete_rows(i)
        elif r['Kode Singkatan'] == 'PINKA' and r['Sudut Pandang'] == 'DURIKOSAMBI':
            assets.delete_rows(i)  # duplicate canonical asset, represented by a Bay below
    for i, r in rows(assets):
        code = r['Kode Singkatan']
        if code == 'MKBRU' and r['Tipe Asset'] == 'Busbar GITET':
            setrow(assets, i, {'Nama Asset / GI': 'GITET Muarakarang Baru', 'No Kerawanan': None})
        elif code == 'GIS MKBRU':
            setrow(assets, i, {'Tipe Asset': 'Busbar GIS', 'Nama Asset / GI': 'GIS Muarakarang Baru'})
        elif code == 'DMGOT':
            setrow(assets, i, {'Tipe Asset': 'Busbar GIS', 'Nama Asset / GI': 'GIS Daan Mogot',
                              'Sudut Pandang': 'DURIKOSAMBI'})
        elif code == 'PINKA':
            setrow(assets, i, {'Nama Asset / GI': 'Pantai Indah Kapuk (PIK)', 'Sudut Pandang': 'MUARAKARANG'})
    existing = {r['Kode Singkatan'] for _, r in rows(assets)}
    for code, hv, lv, unit, view in [
        ('IBT 1 MKBRU', 'GITET_MKBRU', 'GIS MKBRU', 1, 'MUARAKARANG'),
        ('IBT 2 MKBRU', 'GITET_MKBRU', 'GIS MKBRU', 2, 'MUARAKARANG'),
        ('IBT 1 DKSBI', 'GITET_DKSBI', 'DKSBI', 1, 'DURIKOSAMBI')]:
        if code not in existing:
            append(assets, {'Nama Asset / GI': code, 'Kode Singkatan': code, 'Tipe Asset': 'IBT 3-Winding',
                            'No IBT': unit, 'Bus HV': hv, 'Bus LV': lv, 'Sudut Pandang': view,
                            'Tegangan': '500/150 kV'})
    lines = book['Jalur_Transmisi']
    for i, r in rows(lines):
        pair = {r['Dari GI'], r['Ke GI']}
        if pair == {'KBJRK', 'PINKA'}:
            setrow(lines, i, {'Dari GI': 'DMGOT', 'Nama Penghantar': 'SUTT Daan Mogot - PIK',
                             'Sudut Pandang': 'MUARAKARANG;DURIKOSAMBI'})
        elif pair == {'DKSBI', 'DMGOT'}:
            setrow(lines, i, {'Nama Penghantar': 'SUTT Durikosambi - Daan Mogot'})
    if not any({r['Dari GI'], r['Ke GI']} == {'GIS MKBRU', 'MKBRU'} for _, r in rows(lines)):
        append(lines, {'Nama Penghantar': 'Interkoneksi GIS MKBRU - GI MKBRU (incomer)',
                       'Dari GI': 'GIS MKBRU', 'Ke GI': 'MKBRU', 'Tegangan': '150 kV',
                       'Jumlah Sirkit': 2, 'Status Operasi': 'Beroperasi', 'Sudut Pandang': 'MUARAKARANG'})
    if 'Bay' not in book:
        bay = book.create_sheet('Bay')
        bay.append(['Kode GI', 'Nama GI', 'Feeder', 'Jumlah Sirkit', 'Sudut Pandang', 'Tegangan'])
    bay = book['Bay']
    for code, name, feeder, view in [('DMGOT', 'GIS Daan Mogot', 'PINKA', 'MUARAKARANG'),
                                      ('PINKA', 'Pantai Indah Kapuk (PIK)', 'DMGOT', 'DURIKOSAMBI')]:
        if not any(r['Kode GI'] == code for _, r in rows(bay)):
            append(bay, {'Kode GI': code, 'Nama GI': name, 'Feeder': feeder,
                         'Jumlah Sirkit': 2, 'Sudut Pandang': view, 'Tegangan': '150 kV'})
    book['Views'].cell(2, 4, 75)
    book['Views'].cell(3, 4, 76)
    info = book['Info']
    note = 'Koreksi MKBRU dan DMGOT'
    if not any(r[0].value == note for r in info):
        info.append([note, 'Konfirmasi pengguna 2026-09-12; Buku Kerawanan 2026 hal.75-76 (PDF91-92). IBT MKBRU1,2 ke GIS; DKSBI1. DMGOT GIS ke DKSBI dan PIK. Interkoneksi GIS-GI digambar dua sirkit sesuai pasangan pada diagram sumber.'])
    book.save(PATH)


if __name__ == '__main__':
    main()
