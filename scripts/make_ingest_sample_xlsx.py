"""Regenerate samples/ss_cwd_ingest.xlsx from samples/ss_cwd_ingest.json.

The PLN 3-sheet template plus two MANTAPS extensions the base template lacks:
  * an `Info` sheet   -- Kode / Nama / APB subsistem
  * a `Bay` sheet     -- GI drawn as a stub + its feeder (No Kerawanan optional)
  * `No Kerawanan` filled on the asset / line the risk pins to

Run: python scripts/make_ingest_sample_xlsx.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
J = json.loads((ROOT / "samples" / "ss_cwd_ingest.json").read_text(encoding="utf-8"))


def _bold_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True)


def build() -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- Info ----
    ws = wb.create_sheet("Info")
    ws.append(["Kunci", "Nilai"])
    _bold_header(ws)
    ws.append(["Kode Subsistem", J["subsystem"]["code"]])
    ws.append(["Nama Subsistem", J["subsystem"]["name"]])
    ws.append(["APB", J["subsystem"].get("apb", "")])

    # which risk pins to which asset / line / bay / IBT
    pin_ss = {r["pin_key"]: r["seq_no"] for r in J["risks"] if r.get("pin_kind") == "SUBSTATION"}
    pin_ci = {tuple(r["pin_key"].split("-")): r["seq_no"]
              for r in J["risks"] if r.get("pin_kind") == "CIRCUIT" and "-" in r["pin_key"]}
    pin_tx = {r["pin_key"]: r["seq_no"]        # "GITET_CWANG:2" -> seq
              for r in J["risks"] if r.get("pin_kind") == "TRANSFORMER"}

    # ---- Gardu_Induk_dan_Aset ----
    # "Bus 150 kV" names the LV busbar an IBT feeds -- needed when the GITET and
    # its 150 kV bus have different codes (GITET Cawang -> CWBRU, not CWANG).
    ws = wb.create_sheet("Gardu_Induk_dan_Aset")
    ws.append(["No", "Nama Asset / GI", "Kode Singkatan", "Tipe Asset", "Tier (Mulai 0)",
               "Tegangan", "No IBT", "Bus 150 kV", "Status Kerawanan", "No Kerawanan", "Wilayah",
               "Jumlah Trafo", "Jumlah Kapasitor", "Catatan Simbol"])
    _bold_header(ws)
    ibt_lv = {c["from_external_key"]: c["to_external_key"]
              for c in J["connections"]
              if c.get("relation_type") == "IBT_LINK" or c.get("circuit_type_hint") == "IBT_LINK"}
    n = 1
    for o in J["objects"]:
        if o["object_type"] not in ("GITET", "GISTET"):
            continue
        code = o["external_key"].replace("GITET_", "")
        ws.append([n, o["raw_label"], code, "Busbar GITET", 0, "500 kV", None, None,
                   "Normal", None, "Banten"])
        n += 1
    for c in J["connections"]:
        if c.get("relation_type") == "IBT_LINK" or c.get("circuit_type_hint") == "IBT_LINK":
            gk = c["from_external_key"].replace("GITET_", "")
            u = c.get("unit_no", "1")
            lv = c["to_external_key"]
            seq = pin_tx.get(f"{c['from_external_key']}:{u}") or pin_tx.get(f"{gk}:{u}")
            ws.append([n, f"IBT {u} {gk}", f"IBT {u} {gk}", "IBT 3-Winding", 1,
                       "500/150 kV", u, lv, "N-1" if seq else "Normal", seq, "Banten"])
            n += 1
    for o in J["objects"]:
        if o["object_type"] in ("GITET", "GISTET") or o.get("is_bay"):
            continue
        ek = o["external_key"]
        ws.append([n, o["raw_label"], ek, "Busbar GI", o.get("tier_hint"),
                   f'{int(o.get("voltage_hv_kv") or 150)} kV', None, None,
                   "N-1" if ek in pin_ss else "Normal", pin_ss.get(ek), "Banten",
                   o.get("transformer_count", int(bool(o.get("has_transformer")))),
                   o.get("capacitor_count", int(bool(o.get("has_capacitor")))),
                   o.get("symbol_note")])
        n += 1

    # ---- Jalur_Transmisi ----
    ws = wb.create_sheet("Jalur_Transmisi")
    ws.append(["No", "No Kerawanan", "Nama Penghantar", "Dari GI", "Ke GI", "Tegangan",
               "Panjang Saluran (km)", "Jumlah Sirkit", "Status Operasi", "Tingkat Kerawanan",
               "Pembebanan Sirkit 1 (%)", "Pembebanan Sirkit 2 (%)", "Koridor / Wilayah",
               "Tier Dari", "Tier Ke"])
    _bold_header(ws)
    n = 1
    for c in J["connections"]:
        if c.get("relation_type") == "IBT_LINK" or c.get("circuit_type_hint") == "IBT_LINK":
            continue
        fr, to = c["from_external_key"], c["to_external_key"]
        seq = pin_ci.get((fr, to)) or pin_ci.get((to, fr))
        rawan = "Sangat Rawan" if seq else ("Rawan" if c.get("confidence", 1) < 0.7 else "Normal")
        ws.append([n, seq, c.get("note") or f"SUTT {fr} - {to}", fr, to,
                   f'{int((c.get("voltage_kv") or 150))} kV' if c.get("voltage_kv") else "150 kV",
                   None, c.get("circuit_count", 2), "Beroperasi", rawan,
                   None, None, "Banten", None, None])
        n += 1

    # ---- Bay (MANTAPS extension) ----
    ws = wb.create_sheet("Bay")
    ws.append(["No", "Kode GI", "Nama GI", "Feeder (GI Induk)", "Tegangan",
               "Status Operasi", "No Kerawanan"])
    _bold_header(ws)
    n = 1
    for o in J["objects"]:
        if not o.get("is_bay"):
            continue
        ek = o["external_key"]
        ws.append([n, ek, o["raw_label"], o.get("bay_feeder_key"),
                   f'{int(o.get("voltage_hv_kv") or 150)} kV', "Beroperasi",
                   pin_ss.get(ek)])
        n += 1

    # ---- Data_Kerawanan_Detail ----
    ws = wb.create_sheet("Data_Kerawanan_Detail")
    ws.append(["No", "UIT", "Kondisi / Permasalahan", "Dampak", "Mitigasi", "Usulan / Solusi"])
    _bold_header(ws)
    for r in J["risks"]:
        ws.append([r["seq_no"], r.get("uit", "JBB"), r["condition"], r["impact"],
                   r["mitigation"], r["follow_up"]])

    for wsx in wb.worksheets:
        for col in wsx.columns:
            w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            wsx.column_dimensions[col[0].column_letter].width = min(max(w + 2, 10), 55)

    out = ROOT / "samples" / "ss_cwd_ingest.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    p = build()
    from app.services.ingest_parser import parse_upload  # noqa: E402
    d = parse_upload(p.read_bytes(), p.name)
    print(f"wrote {p}")
    print(f"  subsystem: {d['subsystem']}")
    print(f"  {len(d['objects'])} objects, {len(d['connections'])} connections, "
          f"{len(d['risks'])} risks")
    bays = [o for o in d["objects"] if o.get("is_bay")]
    print(f"  bays: {[b['external_key'] + '@' + str(b.get('bay_feeder_key')) for b in bays]}")
    pinned = [r for r in d["risks"] if r.get("pin_key")]
    print(f"  auto-pinned risks: {[(r['seq_no'], r['pin_kind'], r['pin_key']) for r in pinned]}")
