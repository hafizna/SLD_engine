"""Corporate Topology Register (Excel) <-> canonical model.

Excel is an authoring / staging / review format, NOT where the SLD is drawn.
`export_register` writes the canonical model to a multi-sheet workbook;
`import_register` reads that workbook back as ObservedObject / ObservedConnection
(staging) so it goes through reconciliation like any other evidence source.

Sheets:
    00_README            concept notes
    01_SUBSTATION        GI / GITET / GIS / KTT nodes
    02_TRANSFORMER       IBT / trafo + winding summary
    03_GENERATING_UNIT   pembangkit + outlet
    04_CIRCUIT           penghantar / IBT link  (the edges)
    05_SUBSYSTEM         subsystem headers
    06_SS_MEMBERSHIP     GI -> subsystem + role (multi-membership)
    07_RISK              kerawanan (Kondisi/Dampak/Mitigasi/Solusi)
    08_DEFENSE_SCHEME    OLS / OGS / ADS
    09_DS_RELATION       scheme -> object
    10_TOPOLOGY_VERSION  governance
    11_CHANGESET         governance
"""
from __future__ import annotations

from pathlib import Path

from sqlalchemy.orm import Session

from app.models import (
    ChangeSet,
    Circuit,
    DefenseScheme,
    DSRelation,
    GeneratingUnit,
    RiskRecord,
    Substation,
    Subsystem,
    SubsystemMembership,
    TopologyVersion,
    Transformer,
    TransformerWinding,
)

README_ROWS = [
    ("MANTAPS - Corporate Topology Register", ""),
    ("", ""),
    ("Konsep", "Excel = authoring / staging / review topologi relasional. BUKAN tempat menggambar SLD."),
    ("Output", "Engine membentuk canonical model + generate SVG SLD + hitung Tier + overlay."),
    ("Granularitas", "Node = GI/GITET/pembangkit. Edge = penghantar/IBT link. Bay/CB = level PLMS/NMM, tidak wajib."),
    ("Status", "ENERGIZED / NEW_NOT_ENERGIZED / PLANNED / DE_ENERGIZED / OWNED_BY_CUSTOMER (dari warna busbar SLD)."),
    ("Tier", "Dihitung engine per-view dari SOURCE via GI-graph. TIDAK disimpan di GI (satu GI bisa beda Tier per proyeksi)."),
    ("Multi-SS", "Satu GI fisik bisa jadi anggota beberapa subsistem (sheet 06). Objek fisik disimpan sekali."),
    ("Round-trip", "Import kembali -> ObservedObject/ObservedConnection -> reconciliation -> canonical. Tidak menimpa langsung."),
]


def export_register(db: Session, path: str | Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(list(r))
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)

    sheet("00_README", ["Topik", "Keterangan"], README_ROWS)

    # code <-> full-name map for every node kind. The SLD draws the CODE
    # (singkatan); this sheet is the lookup, kept for NMM/CIM alignment.
    _cm_rows = []
    for s in db.query(Substation).order_by(Substation.code).all():
        _cm_rows.append((s.code, s.name, "SUBSTATION", s.substation_type,
                         int(s.voltage_kv), s.apb or "", s.status))
    for t in db.query(Transformer).order_by(Transformer.code).all():
        _cm_rows.append((t.code, t.name, "TRANSFORMER", t.transformer_type, "", "", t.status))
    for g in db.query(GeneratingUnit).order_by(GeneratingUnit.code).all():
        _cm_rows.append((g.code, g.name, "GENERATING_UNIT", g.unit_type or "",
                         int(g.voltage_kv), g.operator or "", g.status))
    sheet("00_CODE_MAP",
          ["Code", "Name", "Kind", "Type", "Voltage_kV", "APB", "Status"],
          _cm_rows)

    sheet("01_SUBSTATION",
          ["Code", "Name", "Type", "Voltage_kV", "APB", "UIT", "Status", "Busbar_Config",
           "Busbar_Note", "Has_Transformer", "Has_Capacitor", "Symbol_Note", "Lat", "Lon",
           "Asset_ID", "Note"],
          [(s.code, s.name, s.substation_type, s.voltage_kv, s.apb, s.uit, s.status,
            s.busbar_config, s.busbar_note, int(s.has_transformer), int(s.has_shunt_capacitor),
            s.symbol_note, s.lat, s.lon, s.asset_id, s.note)
           for s in db.query(Substation).order_by(Substation.code).all()])

    tx_windings: dict[int, list[str]] = {}
    for w in db.query(TransformerWinding).all():
        tx_windings.setdefault(w.transformer_id, []).append(f"{w.winding_no}:{w.voltage_kv}kV({w.role})")
    sheet("02_TRANSFORMER",
          ["Code", "Name", "Type", "Substation_Code", "Unit_No", "Rating_MVA", "Winding_Count",
           "Windings", "Status", "Asset_ID", "Note"],
          [(t.code, t.name, t.transformer_type,
            db.get(Substation, t.substation_id).code if t.substation_id else "",
            t.unit_no, t.rating_mva, t.winding_count,
            " | ".join(tx_windings.get(t.id, [])), t.status, t.asset_id, t.note)
           for t in db.query(Transformer).order_by(Transformer.code).all()])

    sheet("03_GENERATING_UNIT",
          ["Code", "Name", "Unit_Type", "Voltage_kV", "Rated_MW", "Unit_Count",
           "Outlet_Substation_Code", "Operator", "Status", "Note"],
          [(g.code, g.name, g.unit_type, g.voltage_kv, g.rated_mw, g.unit_count,
            db.get(Substation, g.outlet_substation_id).code if g.outlet_substation_id else "",
            g.operator, g.status, g.note)
           for g in db.query(GeneratingUnit).order_by(GeneratingUnit.code).all()])

    sub_code = {s.id: s.code for s in db.query(Substation).all()}
    sheet("04_CIRCUIT",
          ["Code", "Name", "Circuit_Type", "Voltage_kV", "From_Substation", "To_Substation",
           "Circuit_Count", "Single_Phi", "Status", "Scenario", "Confidence", "Note"],
          [(c.code, c.name, c.circuit_type, c.voltage_kv,
            sub_code.get(c.from_substation_id, "?"), sub_code.get(c.to_substation_id, "?"),
            c.circuit_count, int(c.single_phi), c.status, c.scenario_id, c.confidence, c.note)
           for c in db.query(Circuit).order_by(Circuit.code).all()])

    sheet("05_SUBSYSTEM",
          ["Code", "Name", "APB", "Source_Ref"],
          [(s.code, s.name, s.apb, s.source_ref)
           for s in db.query(Subsystem).order_by(Subsystem.code).all()])

    ss_code = {s.id: s.code for s in db.query(Subsystem).all()}
    sheet("06_SS_MEMBERSHIP",
          ["Subsystem", "Node_Kind", "Node_Code", "Role", "External_Subsystem", "Display_Order"],
          [(ss_code.get(m.subsystem_id, "?"), m.node_kind,
            sub_code.get(m.node_id, str(m.node_id)) if m.node_kind == "SUBSTATION" else str(m.node_id),
            m.role, m.external_subsystem, m.display_order)
           for m in db.query(SubsystemMembership).all()])

    sheet("07_RISK",
          ["Risk_Key", "Subsystem", "Seq_No", "UIT", "Attach_Kind", "Attach_Label", "Title",
           "Kondisi", "Dampak", "Mitigasi", "Usulan_Solusi", "Priority", "Status"],
          [(r.risk_key, ss_code.get(r.subsystem_id, ""), r.seq_no, r.uit, r.attach_kind,
            r.attach_label, r.title, r.condition, r.impact, r.mitigation, r.follow_up,
            r.priority, r.status)
           for r in db.query(RiskRecord).order_by(RiskRecord.seq_no).all()])

    sheet("08_DEFENSE_SCHEME",
          ["Scheme_Key", "Name", "Type", "Status", "Scope", "Target_MW", "Target_MW_Planned",
           "Stages", "Reference", "Note"],
          [(d.scheme_key, d.name, d.scheme_type, d.status, d.scope_level, d.target_mw,
            d.target_mw_planned, d.stages, d.reference, d.note)
           for d in db.query(DefenseScheme).order_by(DefenseScheme.scheme_key).all()])

    scheme_key = {d.id: d.scheme_key for d in db.query(DefenseScheme).all()}
    sheet("09_DS_RELATION",
          ["Scheme_Key", "Subsystem", "Attach_Kind", "Attach_Label", "Role", "Coverage_Note"],
          [(scheme_key.get(rel.scheme_id, "?"), ss_code.get(rel.subsystem_id, ""),
            rel.attach_kind, rel.attach_label, rel.role, rel.coverage_note)
           for rel in db.query(DSRelation).all()])

    sheet("10_TOPOLOGY_VERSION",
          ["Version_Key", "Status", "Effective_Date", "Reviewer", "Approver", "Description"],
          [(v.version_key, v.status, v.effective_date, v.reviewer, v.approver, v.description)
           for v in db.query(TopologyVersion).order_by(TopologyVersion.version_key).all()])

    ver_key = {v.id: v.version_key for v in db.query(TopologyVersion).all()}
    sheet("11_CHANGESET",
          ["Change_Key", "Version", "Action", "Target_Kind", "Target_Ref", "Status", "Description"],
          [(cs.change_key, ver_key.get(cs.version_id, ""), cs.action, cs.target_kind,
            cs.target_ref, cs.status, cs.description)
           for cs in db.query(ChangeSet).order_by(ChangeSet.change_key).all()])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def import_register(db: Session, path: str | Path):
    """Read 01_SUBSTATION + 04_CIRCUIT back as a staged observation batch.

    Returns the SourceDocument. Endpoints are matched by substation code; the
    reconciliation engine still decides AUTO_MATCH / REVIEW / CREATE_NEW.
    """
    import openpyxl

    from app.models import ObservedConnection, ObservedObject, SourceDocument
    from app.services.reconciliation import normalize_name

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    doc = SourceDocument(
        filename=Path(path).name,
        document_type="EXCEL_REGISTER",
        source_ref="Corporate Topology Register import",
    )
    db.add(doc)
    db.flush()

    obs_by_code: dict[str, ObservedObject] = {}
    ws = wb["01_SUBSTATION"]
    header = [c.value for c in next(ws.iter_rows())]
    idx = {name: i for i, name in enumerate(header)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx["Code"]]:
            continue
        code = str(row[idx["Code"]]).strip()
        o = ObservedObject(
            document_id=doc.id, external_key=code, object_type="SUBSTATION",
            raw_label=str(row[idx["Name"]] or code),
            normalized_name=normalize_name(str(row[idx["Name"]] or code)),
            site_name=str(row[idx["Name"]] or code),
            voltage_hv_kv=row[idx["Voltage_kV"]],
            status_hint=row[idx["Status"]],
            confidence=1.0,
        )
        db.add(o)
        db.flush()
        obs_by_code[code] = o

    if "04_CIRCUIT" in wb.sheetnames:
        ws = wb["04_CIRCUIT"]
        header = [c.value for c in next(ws.iter_rows())]
        idx = {name: i for i, name in enumerate(header)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[idx["Code"]]:
                continue
            a = obs_by_code.get(str(row[idx["From_Substation"]]).strip())
            b = obs_by_code.get(str(row[idx["To_Substation"]]).strip())
            if not a or not b:
                continue
            db.add(ObservedConnection(
                document_id=doc.id, from_observed_id=a.id, to_observed_id=b.id,
                relation_type="CONNECTED_TO",
                circuit_type_hint=row[idx["Circuit_Type"]],
                status_hint=row[idx["Status"]],
                confidence=float(row[idx["Confidence"]] or 1.0),
            ))

    db.commit()
    return doc
